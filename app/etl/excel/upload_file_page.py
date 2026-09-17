"""ETL: Excel — Upload File (session, upload, analyze mapping table)."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QShowEvent
from PySide6.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QPushButton,
    QSizePolicy,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.etl.widgets.etl_connection_picker import EtlConnectionHeaderPicker
from core.api import (
    api_analyze_import_file,
    api_analyze_import_if_present,
    api_get_all_import_session_ids,
    api_save_import_mapping,
    api_upload_import_file,
    import_sql_datatypes_for_db,
    import_target_data_type_to_java,
    import_target_data_type_to_sql,
)
from core.etl_connection_context import etl_connection_id, get_etl_connection_context
from core.user_context import get_user_profile
from ui.auto_hide_message import cancel_auto_hide_message, show_auto_hiding_message
from ui.data_table import apply_data_table_appearance, attach_table_copy_shortcut
from ui.form_combobox_style import apply_form_combobox_field, install_combo_popup_below_field
from ui.form_page_styles import (
    FORM_ERROR_LABEL_STYLE,
    FORM_INPUT_STYLE,
    FORM_LABEL_STYLE,
    FORM_PRIMARY_BUTTON_STYLESHEET,
    FORM_SINGLELINE_FIELD_HEIGHT_PX,
    LIST_PAGE_HEADER_HEIGHT_PX,
    LIST_PAGE_HEADER_LAYOUT_MARGINS,
    LIST_PAGE_HEADER_LAYOUT_SPACING,
    LIST_PAGE_HEADER_STYLESHEET,
)
from ui.widgets.required_label import field_caption_label, labeled_field_block
from ui.theme import Theme

_HEADER_BTN_DISABLED_STYLESHEET = (
    f"QPushButton {{ background: {Theme.HEADER_ACCENT}; color: {Theme.PANEL_TEXT_BRIGHT}; border: none; "
    "border-radius: 6px; padding: 5px 16px; font-size: 12px; font-weight: 500; }"
    f"QPushButton:hover:!disabled {{ background: {Theme.HEADER_ACCENT_HOVER}; }}"
    f"QPushButton:pressed:!disabled {{ background: {Theme.HEADER_ACCENT_PRESSED}; }}"
    f"QPushButton:disabled {{ background: #1e293b; color: {Theme.TEXT_SECONDARY}; }}"
)

_UPLOAD_FILE_FILTER = (
    "Spreadsheets (*.xlsx *.xls *.csv);;Excel (*.xlsx *.xls);;CSV (*.csv);;"
    "JSON (*.json);;All files (*.*)"
)

# (label, codepage value sent to API)
_CODEPAGE_OPTIONS: tuple[tuple[str, str], ...] = (
    ("UTF-8 (Unicode) — 65001", "65001"),
    ("Windows Arabic — 1256", "1256"),
    ("ISO-8859-6 Arabic — 28596", "28596"),
    ("Windows Western European — 1252", "1252"),
    ("Windows Cyrillic — 1251", "1251"),
    ("Windows Central European — 1250", "1250"),
    ("Windows Greek — 1253", "1253"),
    ("Windows Turkish — 1254", "1254"),
    ("Windows Hebrew — 1255", "1255"),
    ("Shift-JIS Japanese — 932", "932"),
    ("GBK Chinese Simplified — 936", "936"),
    ("Big5 Chinese Traditional — 950", "950"),
    ("EUC-KR Korean — 949", "949"),
)


def _text_looks_mojibake_or_question_marks(text: str) -> bool:
    s = str(text or "")
    if not s:
        return False
    q = s.count("?")
    if q >= 2 and (q / max(len(s), 1)) >= 0.35:
        return True
    return "\ufffd" in s


def _read_local_workbook_headers(path: str) -> list[dict[str, Any]]:
    """Read sheet titles + header row via openpyxl (Unicode-safe for .xlsx)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    p = Path(path)
    if not p.is_file():
        return []
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return []
    out: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            headers: list[str] = []
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                for cell in row:
                    headers.append("" if cell is None else str(cell))
                break
            out.append({"name": str(ws.title or ""), "headers": headers})
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return out


def _enrich_analyze_sheets_unicode(
    sheets: list[Any],
    *,
    file_path: str,
    multi_language: bool,
) -> list[Any]:
    """Replace corrupted ``?`` sheet/column names using the local workbook when possible."""
    if not multi_language or not sheets or not file_path:
        return sheets
    local = _read_local_workbook_headers(file_path)
    if not local:
        return sheets
    enriched: list[Any] = []
    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict):
            enriched.append(sheet)
            continue
        sheet_copy = deepcopy(sheet)
        local_sheet = local[index] if index < len(local) else None
        if local_sheet:
            api_name = str(sheet_copy.get("name") or "")
            local_name = str(local_sheet.get("name") or "")
            if local_name and (
                not api_name
                or _text_looks_mojibake_or_question_marks(api_name)
                or ("?" in api_name and "?" not in local_name)
            ):
                sheet_copy["name"] = local_name
            columns = sheet_copy.get("columns")
            local_headers = local_sheet.get("headers") if isinstance(local_sheet, dict) else None
            if isinstance(columns, list) and isinstance(local_headers, list):
                new_cols: list[Any] = []
                for col_i, col in enumerate(columns):
                    if not isinstance(col, dict):
                        new_cols.append(col)
                        continue
                    col_copy = dict(col)
                    local_hdr = (
                        str(local_headers[col_i]) if col_i < len(local_headers) else ""
                    )
                    if local_hdr:
                        name = str(col_copy.get("name") or "")
                        tgt = str(col_copy.get("tgtColumnName") or "")
                        if not name or _text_looks_mojibake_or_question_marks(name):
                            col_copy["name"] = local_hdr
                        if not tgt or _text_looks_mojibake_or_question_marks(tgt):
                            col_copy["tgtColumnName"] = local_hdr
                    new_cols.append(col_copy)
                sheet_copy["columns"] = new_cols
        enriched.append(sheet_copy)
    return enriched


_ANALYZE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("name", "Name"),
    ("ordinal", "Ordinal"),
    ("tgtColumnName", "TGT Column Name"),
    ("tgtDatatype", "TGT Datatype"),
    ("length", "Length"),
    ("decimal", "Decimal"),
    ("precision", "Precision"),
)

_TGT_DATATYPE_OPTIONS = (
    "NVARCHAR",
    "VARCHAR",
    "INT",
    "BIGINT",
    "DECIMAL",
    "FLOAT",
    "BIT",
    "DATE",
    "DATETIME",
    "DATETIME2",
    "TEXT",
)

_SHEET_TAB_STYLE = (
    "QPushButton { background: transparent; color: #334155; border: none; "
    "border-bottom: 2px solid transparent; padding: 6px 16px; font-size: 11px; }"
    "QPushButton:hover { background: #e2e8f0; color: #0f172a; }"
    "QPushButton:checked { background: #ffffff; color: #1d4ed8; "
    "border-bottom: 3px solid #2563eb; font-weight: 700; }"
)

_SHEET_TAB_SAVED_STYLE = (
    "QPushButton { background: transparent; color: #15803d; border: none; "
    "border-bottom: 2px solid transparent; padding: 6px 16px; font-size: 11px; font-weight: 700; }"
    "QPushButton:hover { background: #dcfce7; color: #166534; }"
    "QPushButton:checked { background: #ffffff; color: #15803d; "
    "border-bottom: 3px solid #22c55e; font-weight: 700; }"
)

_CHOOSE_SHEET_STYLE = (
    "QLabel#chooseSheetChip { background: #eff6ff; color: #1d4ed8; border: 1px solid #bfdbfe; "
    "border-radius: 10px; padding: 3px 10px; font-size: 11px; font-weight: 700; }"
)

_UPLOAD_CARD_STYLE = """
    QFrame#uploadCard {
        background: #ffffff;
        border: 1px solid #dbe3ec;
        border-radius: 10px;
    }
    QLabel#uploadCardTitle {
        color: #0f172a;
        font-size: 15px;
        font-weight: 700;
        background: transparent;
        border: none;
    }
    QLabel#uploadCardHint {
        color: #64748b;
        font-size: 11px;
        background: transparent;
        border: none;
    }
"""

_TARGET_TABLE_CARD_STYLE = f"""
    QFrame#targetTableCard {{
        background: {Theme.BG_WHITE};
        border: 1px solid {Theme.BORDER_DEFAULT};
        border-radius: 10px;
    }}
"""

_WORKBOOK_STYLE = """
    QFrame#workbookCard {
        background: #ffffff;
        border: 1px solid #cbd5e1;
        border-radius: 8px;
    }
    QFrame#workbookToolbar {
        background: #1e3a5f;
        border: none;
        border-top-left-radius: 7px;
        border-top-right-radius: 7px;
    }
    QLabel#workbookTitle {
        color: #ffffff;
        font-size: 13px;
        font-weight: 700;
        background: transparent;
        border: none;
    }
    QLabel#workbookStatus {
        color: #bfdbfe;
        font-size: 10px;
        background: transparent;
        border: none;
    }
    QFrame#sheetBar {
        background: #f1f5f9;
        border: none;
        border-top: 1px solid #cbd5e1;
    }
"""

_EXCEL_TABLE_STYLE = """
    QTableWidget {
        background: #ffffff;
        alternate-background-color: #ffffff;
        color: #1f2937;
        border: none;
        gridline-color: #cbd5e1;
        font-size: 11px;
        selection-background-color: #dbeafe;
        selection-color: #0f172a;
    }
    QTableWidget::item {
        padding: 3px 7px;
        border: none;
    }
    QTableWidget::item:selected {
        background: #dbeafe;
        color: #0f172a;
    }
    QHeaderView::section {
        background: #f3f4f6;
        color: #1f2937;
        border: none;
        border-right: 1px solid #cbd5e1;
        border-bottom: 1px solid #94a3b8;
        padding: 5px 7px;
        font-size: 11px;
        font-weight: 700;
    }
    QTableCornerButton::section {
        background: #e5e7eb;
        border: none;
        border-right: 1px solid #94a3b8;
        border-bottom: 1px solid #94a3b8;
    }
"""


class UploadFilePage(QWidget):
    """Upload File: session select, file upload, then analyze mapping grid."""

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._loaded_once = False
        self._selected_path = ""
        self._uploading = False
        self._analyzing = False
        self._saving_mapping = False
        self._edit_mode = False
        self._sheets: list[dict[str, Any]] = []
        self._active_sheet_index = 0
        self._sheet_tab_buttons: list[QPushButton] = []
        self._saved_sheet_keys: set[str] = set()
        self._ctx = get_etl_connection_context()
        self._ctx.connection_selected.connect(self._on_connection_selected)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        header = QWidget()
        header.setStyleSheet(LIST_PAGE_HEADER_STYLESHEET)
        header.setFixedHeight(LIST_PAGE_HEADER_HEIGHT_PX)
        hl = QHBoxLayout(header)
        hl.setContentsMargins(*LIST_PAGE_HEADER_LAYOUT_MARGINS)
        hl.setSpacing(LIST_PAGE_HEADER_LAYOUT_SPACING)
        hl.addWidget(QLabel("Excel: Upload File"))
        hl.addStretch()
        self.connection_picker = EtlConnectionHeaderPicker()
        hl.addWidget(self.connection_picker)
        self.upload_btn = QPushButton("Upload")
        self.upload_btn.setFixedWidth(100)
        self.upload_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.upload_btn.setStyleSheet(_HEADER_BTN_DISABLED_STYLESHEET)
        self.upload_btn.clicked.connect(self._handle_upload)
        hl.addWidget(self.upload_btn)
        root.addWidget(header)

        body = QWidget()
        body.setStyleSheet(f"background: {Theme.BG_PAGE_ALT};")
        self._body_layout = QVBoxLayout(body)
        self._body_layout.setContentsMargins(24, 18, 24, 18)
        self._body_layout.setSpacing(14)

        top_row = QWidget()
        top_row_layout = QHBoxLayout(top_row)
        top_row_layout.setContentsMargins(0, 0, 0, 0)
        top_row_layout.setSpacing(16)
        top_row_layout.setAlignment(Qt.AlignmentFlag.AlignTop)

        form = QFrame()
        form.setObjectName("uploadCard")
        form.setStyleSheet(_UPLOAD_CARD_STYLE)
        form.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        form.setMaximumWidth(840)
        form_layout = QVBoxLayout(form)
        form_layout.setContentsMargins(20, 16, 20, 18)
        form_layout.setSpacing(10)

        card_title = QLabel("Upload a source file")
        card_title.setObjectName("uploadCardTitle")
        form_layout.addWidget(card_title)
        card_hint = QLabel(
            "Select an import session and choose a supported workbook. "
            "After upload, its sheets and columns are analyzed automatically."
        )
        card_hint.setObjectName("uploadCardHint")
        card_hint.setWordWrap(True)
        form_layout.addWidget(card_hint)
        form_layout.addSpacing(4)

        field_w = 520

        session_lbl = field_caption_label("Session ID*", FORM_LABEL_STYLE)
        self.session_combo = QComboBox()
        self.session_combo.setPlaceholderText("Select session name")
        apply_form_combobox_field(
            self.session_combo,
            height_px=FORM_SINGLELINE_FIELD_HEIGHT_PX,
            min_width=field_w,
        )
        install_combo_popup_below_field(self.session_combo)
        self.session_combo.setCurrentIndex(-1)
        self.session_combo.currentIndexChanged.connect(self._on_session_changed)
        form_layout.addWidget(labeled_field_block(session_lbl, self.session_combo))

        file_lbl = field_caption_label("File Upload*", FORM_LABEL_STYLE)
        file_row = QWidget()
        file_row_layout = QHBoxLayout(file_row)
        file_row_layout.setContentsMargins(0, 0, 0, 0)
        file_row_layout.setSpacing(8)
        self.file_path_edit = QLineEdit()
        self.file_path_edit.setReadOnly(True)
        self.file_path_edit.setPlaceholderText("No file selected")
        self.file_path_edit.setStyleSheet(FORM_INPUT_STYLE)
        self.file_path_edit.setFixedHeight(FORM_SINGLELINE_FIELD_HEIGHT_PX)
        self.file_path_edit.setMinimumWidth(field_w - 108)
        self.file_path_edit.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        browse_btn = QPushButton("Browse…")
        browse_btn.setFixedWidth(100)
        browse_btn.setFixedHeight(FORM_SINGLELINE_FIELD_HEIGHT_PX)
        browse_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        browse_btn.setStyleSheet(
            "QPushButton { background: #f8fafc; color: #334155; border: 1px solid #cbd5e1; "
            "border-radius: 4px; padding: 4px 12px; font-size: 11px; font-weight: 600; }"
            "QPushButton:hover { background: #eef2f7; border-color: #94a3b8; }"
        )
        browse_btn.clicked.connect(self._browse_file)
        file_row_layout.addWidget(self.file_path_edit, 1)
        file_row_layout.addWidget(browse_btn)
        form_layout.addWidget(labeled_field_block(file_lbl, file_row))

        self.multi_language_check = QCheckBox("Enable multi-language (Unicode)")
        self.multi_language_check.setChecked(True)
        self.multi_language_check.setCursor(Qt.CursorShape.PointingHandCursor)
        self.multi_language_check.setStyleSheet(
            "QCheckBox { color: #334155; font-size: 11px; font-weight: 600; }"
            "QCheckBox::indicator { width: 14px; height: 14px; }"
        )
        self.multi_language_check.setToolTip(
            "Preserve Arabic and other non-English sheet/column names. "
            "Also repairs names locally from the workbook when the server returns '?'."
        )
        form_layout.addWidget(self.multi_language_check)

        codepage_lbl = field_caption_label("Codepage", FORM_LABEL_STYLE)
        self.codepage_combo = QComboBox()
        for label, value in _CODEPAGE_OPTIONS:
            self.codepage_combo.addItem(label, value)
        self.codepage_combo.setCurrentIndex(0)  # UTF-8
        apply_form_combobox_field(
            self.codepage_combo,
            height_px=FORM_SINGLELINE_FIELD_HEIGHT_PX,
            min_width=field_w,
        )
        install_combo_popup_below_field(self.codepage_combo)
        self.codepage_combo.setToolTip(
            "Character encoding / Windows code page used when reading the file. "
            "Use UTF-8 for .xlsx, or Windows Arabic (1256) for older .xls/CSV Arabic files."
        )
        form_layout.addWidget(labeled_field_block(codepage_lbl, self.codepage_combo))
        self.multi_language_check.toggled.connect(self.codepage_combo.setEnabled)
        self.codepage_combo.setEnabled(True)

        self.message_label = QLabel()
        self.message_label.setStyleSheet(FORM_ERROR_LABEL_STYLE)
        self.message_label.setWordWrap(True)
        self.message_label.setTextInteractionFlags(
            Qt.TextInteractionFlag.TextSelectableByMouse
            | Qt.TextInteractionFlag.TextSelectableByKeyboard
        )
        self.message_label.setVisible(False)
        form_layout.addWidget(self.message_label)

        top_row_layout.addWidget(form, 1, Qt.AlignmentFlag.AlignTop)

        target_card = QFrame()
        target_card.setObjectName("targetTableCard")
        target_card.setStyleSheet(_TARGET_TABLE_CARD_STYLE)
        target_card.setFixedWidth(320)
        target_card.setSizePolicy(QSizePolicy.Policy.Fixed, QSizePolicy.Policy.Fixed)
        target_layout = QVBoxLayout(target_card)
        target_layout.setContentsMargins(16, 14, 16, 16)
        target_layout.setSpacing(6)
        target_lbl = field_caption_label("Target table name", FORM_LABEL_STYLE)
        self.target_table_edit = QLineEdit()
        self.target_table_edit.setPlaceholderText("Enter target table name")
        self.target_table_edit.setStyleSheet(FORM_INPUT_STYLE)
        self.target_table_edit.setFixedHeight(FORM_SINGLELINE_FIELD_HEIGHT_PX)
        target_layout.addWidget(labeled_field_block(target_lbl, self.target_table_edit))
        top_row_layout.addWidget(
            target_card,
            0,
            Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignRight,
        )

        self._body_layout.addWidget(top_row, 0, Qt.AlignmentFlag.AlignTop)
        self.analyze_panel = QFrame()
        self.analyze_panel.setObjectName("workbookCard")
        self.analyze_panel.setStyleSheet(_WORKBOOK_STYLE)
        self.analyze_panel.setVisible(False)
        self.analyze_panel.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        analyze_layout = QVBoxLayout(self.analyze_panel)
        analyze_layout.setContentsMargins(0, 0, 0, 0)
        analyze_layout.setSpacing(0)

        workbook_toolbar = QFrame()
        workbook_toolbar.setObjectName("workbookToolbar")
        workbook_toolbar.setFixedHeight(44)
        workbook_toolbar_layout = QHBoxLayout(workbook_toolbar)
        workbook_toolbar_layout.setContentsMargins(14, 0, 12, 0)
        workbook_toolbar_layout.setSpacing(10)
        workbook_title = QLabel("Workbook analysis")
        workbook_title.setObjectName("workbookTitle")
        workbook_toolbar_layout.addWidget(workbook_title)
        workbook_status = QLabel("Ready to review")
        workbook_status.setObjectName("workbookStatus")
        workbook_toolbar_layout.addWidget(workbook_status)
        workbook_toolbar_layout.addStretch()
        self.choose_sheet_chip = QLabel("")
        self.choose_sheet_chip.setObjectName("chooseSheetChip")
        self.choose_sheet_chip.setStyleSheet(_CHOOSE_SHEET_STYLE)
        workbook_toolbar_layout.addWidget(self.choose_sheet_chip)
        self.edit_btn = QPushButton("Edit mapping")
        self.edit_btn.setFixedWidth(110)
        self.edit_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.edit_btn.setStyleSheet(
            "QPushButton { background: #ffffff; color: #1e3a5f; border: none; border-radius: 5px; "
            "padding: 5px 12px; font-size: 11px; font-weight: 700; }"
            "QPushButton:hover { background: #eff6ff; }"
        )
        self.edit_btn.clicked.connect(self._toggle_edit_mode)
        workbook_toolbar_layout.addWidget(self.edit_btn)
        self.save_btn = QPushButton("Save mapping")
        self.save_btn.setFixedWidth(110)
        self.save_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.save_btn.setStyleSheet(
            "QPushButton { background: #0f2340; color: #ffffff; border: 1px solid #60a5fa; "
            "border-radius: 5px; padding: 5px 12px; font-size: 11px; font-weight: 700; }"
            "QPushButton:hover { background: #172f52; }"
        )
        self.save_btn.clicked.connect(self._handle_save)
        workbook_toolbar_layout.addWidget(self.save_btn)
        analyze_layout.addWidget(workbook_toolbar)

        tabs_row = QHBoxLayout()
        tabs_row.setSpacing(8)
        self.sheet_tabs_host = QWidget()
        self.sheet_tabs_layout = QHBoxLayout(self.sheet_tabs_host)
        self.sheet_tabs_layout.setContentsMargins(8, 0, 0, 0)
        self.sheet_tabs_layout.setSpacing(2)
        self.sheet_tabs_layout.addStretch()
        tabs_row.addWidget(self.sheet_tabs_host, 1)

        self.analyze_table = QTableWidget()
        self.analyze_table.setColumnCount(len(_ANALYZE_COLUMNS))
        self.analyze_table.setHorizontalHeaderLabels([label for _, label in _ANALYZE_COLUMNS])
        apply_data_table_appearance(
            self.analyze_table,
            read_only=True,
            stretch_last_section=False,
            hide_vertical_header=False,
            sort_indicator_shown=False,
            alternating_row_colors=False,
        )
        self.analyze_table.setStyleSheet(_EXCEL_TABLE_STYLE)
        self.analyze_table.setShowGrid(True)
        self.analyze_table.setCornerButtonEnabled(True)
        self.analyze_table.setMinimumHeight(260)
        self.analyze_table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.analyze_table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        hh = self.analyze_table.horizontalHeader()
        hh.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        hh.setStretchLastSection(True)
        hh.setDefaultAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        vh = self.analyze_table.verticalHeader()
        vh.setDefaultSectionSize(28)
        vh.setMinimumWidth(42)
        vh.setDefaultAlignment(Qt.AlignmentFlag.AlignCenter)
        attach_table_copy_shortcut(self.analyze_table)
        analyze_layout.addWidget(self.analyze_table, 1)

        sheet_bar = QFrame()
        sheet_bar.setObjectName("sheetBar")
        sheet_bar.setFixedHeight(36)
        sheet_bar_layout = QHBoxLayout(sheet_bar)
        sheet_bar_layout.setContentsMargins(8, 0, 8, 0)
        sheet_bar_layout.setSpacing(8)
        sheet_icon = QLabel("▦")
        sheet_icon.setStyleSheet(
            "color: #2563eb; font-size: 15px; font-weight: 700; background: transparent;"
        )
        sheet_bar_layout.addWidget(sheet_icon)
        sheet_bar_layout.addLayout(tabs_row, 1)
        analyze_layout.addWidget(sheet_bar)

        self._body_layout.addWidget(self.analyze_panel, 1)
        root.addWidget(body, 1)

        footer = QWidget()
        footer.setStyleSheet(f"background: {Theme.BG_WHITE}; border-top: 1px solid {Theme.BORDER_DEFAULT};")
        footer.setFixedHeight(56)
        footer_layout = QHBoxLayout(footer)
        footer_layout.setContentsMargins(24, 10, 24, 10)
        footer_layout.setSpacing(12)
        self.footer_save_btn = QPushButton("Save")
        self.footer_save_btn.setFixedWidth(110)
        self.footer_save_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.footer_save_btn.setStyleSheet(FORM_PRIMARY_BUTTON_STYLESHEET)
        self.footer_save_btn.clicked.connect(self._handle_save)
        footer_layout.addWidget(self.footer_save_btn, 0, Qt.AlignmentFlag.AlignLeft)
        footer_layout.addStretch(1)
        root.addWidget(footer)
    def showEvent(self, event: QShowEvent) -> None:
        super().showEvent(event)
        self._refresh_connections()
        if not self._loaded_once:
            self._loaded_once = True
            self._load_sessions()

    def go_to_nav_item(self, item_name: str) -> None:
        _ = item_name
        self._refresh_connections()
        self._load_sessions()

    def selected_session_id(self) -> str:
        return str(self.session_combo.currentData() or "").strip()

    def selected_session_name(self) -> str:
        return str(self.session_combo.currentText() or "").strip()

    def selected_file_path(self) -> str:
        return self._selected_path

    def selected_connection(self) -> dict[str, Any] | None:
        return self._ctx.current_connection()

    def selected_target_table_name(self) -> str:
        return self.target_table_edit.text().strip()

    def _selected_db_type(self) -> str:
        conn = self.selected_connection() or {}
        return str(conn.get("dbType") or conn.get("databaseType") or conn.get("db_type") or "").strip()

    def _datatype_options(self) -> tuple[str, ...]:
        options = import_sql_datatypes_for_db(self._selected_db_type())
        return options if options else _TGT_DATATYPE_OPTIONS

    def _sheet_key(self, sheet: dict[str, Any] | None, index: int | None = None) -> str:
        idx = self._active_sheet_index if index is None else index
        sheet_id = self._resolve_sheet_id(sheet or {})
        if sheet_id is not None:
            return f"id:{sheet_id}"
        name = str((sheet or {}).get("name") or f"Sheet{idx + 1}").strip()
        return f"name:{name.lower()}"

    def _on_connection_selected(self, _conn: object) -> None:
        if not self._sheets:
            return
        if self._edit_mode:
            self._capture_table_edits_into_sheet()
        self._apply_sql_datatypes_for_connection()
        self._render_active_sheet()

    def _apply_sql_datatypes_for_connection(self) -> None:
        db_type = self._selected_db_type()
        for sheet in self._sheets:
            columns = sheet.get("columns") if isinstance(sheet, dict) else None
            if not isinstance(columns, list):
                continue
            for col in columns:
                if not isinstance(col, dict):
                    continue
                current = col.get("tgtDatatype") or col.get("datatype") or ""
                col["tgtDatatype"] = import_target_data_type_to_sql(current, db_type)

    def _token(self) -> str | None:
        profile = get_user_profile() or {}
        return profile.get("token") or profile.get("accessToken")

    def _refresh_connections(self) -> None:
        result = self._ctx.refresh_connections(self._token())
        if not result.get("success"):
            self._show_error(str(result.get("message") or "Failed to load connections."))

    def _clear_message(self) -> None:
        cancel_auto_hide_message(self, self.message_label)
        self.message_label.setText("")
        self.message_label.setVisible(False)

    def _show_error(self, message: str) -> None:
        show_auto_hiding_message(self, self.message_label, message, error=True)

    def _show_success(self, message: str) -> None:
        show_auto_hiding_message(self, self.message_label, message, error=False)

    def _browse_file(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Select file to upload",
            "",
            _UPLOAD_FILE_FILTER,
        )
        if not path:
            return
        self._selected_path = path
        self.file_path_edit.setText(Path(path).name)
        self._clear_message()

    def _load_sessions(self) -> None:
        previous_id = self.selected_session_id()
        self.upload_btn.setEnabled(False)
        try:
            result = api_get_all_import_session_ids(self._token())
        finally:
            self.upload_btn.setEnabled(True)

        self.session_combo.blockSignals(True)
        self.session_combo.clear()
        sessions = result.get("data") if result.get("success") else []
        if not isinstance(sessions, list):
            sessions = []
        for row in sessions:
            if not isinstance(row, dict):
                continue
            sid = str(row.get("sessionId") or "").strip()
            name = str(row.get("sessionName") or "").strip()
            if not name and not sid:
                continue
            self.session_combo.addItem(name or sid, sid)

        if previous_id:
            idx = self.session_combo.findData(previous_id)
            self.session_combo.setCurrentIndex(idx if idx >= 0 else -1)
        else:
            self.session_combo.setCurrentIndex(-1)
        self.session_combo.blockSignals(False)
        if self.selected_session_id():
            self._check_analyze_if_present(self.selected_session_id())

    def _on_session_changed(self, _index: int = -1) -> None:
        session_id = self.selected_session_id()
        if not session_id:
            self._clear_analyze_panel()
            self._clear_session_target_fields()
            return
        self._check_analyze_if_present(session_id)

    def _clear_session_target_fields(self) -> None:
        self.target_table_edit.clear()

    def _resolve_connection_id(self, value: Any) -> Any:
        if isinstance(value, dict):
            for key in ("id", "connectionId", "connectionID", "connection_id"):
                if value.get(key) is not None and str(value.get(key)).strip() != "":
                    return value.get(key)
            return None
        return value

    def _sheet_target_table_name(self, sheet: dict[str, Any] | None) -> str:
        if not isinstance(sheet, dict):
            return ""
        for key in ("targetTableName", "tableName"):
            value = sheet.get(key)
            if value is not None and str(value).strip():
                return str(value).strip()
        raw = sheet.get("_raw") if isinstance(sheet.get("_raw"), dict) else {}
        for key in (
            "targetTableName",
            "target_table_name",
            "tgtTableName",
            "targetTable",
            "tableName",
        ):
            value = raw.get(key)
            if value is not None and str(value).strip():
                return str(value).strip()
        return ""

    def _apply_session_target_meta(self, result: dict[str, Any], sheets: list[Any]) -> None:
        connection_id = self._resolve_connection_id(result.get("targetConnectionId"))
        raw = result.get("raw") if isinstance(result.get("raw"), dict) else {}
        if connection_id is None and raw:
            connection_id = self._resolve_connection_id(
                raw.get("connectionId") or raw.get("connection") or raw.get("targetConnectionId")
            )

        table_name = str(result.get("targetTableName") or "").strip()
        # Prefer the first sheet that already has a mapped table name.
        if not table_name:
            for sheet in sheets:
                if not isinstance(sheet, dict):
                    continue
                table_name = self._sheet_target_table_name(sheet)
                if table_name:
                    break

        if table_name:
            self.target_table_edit.setText(table_name)
        else:
            self.target_table_edit.clear()

        if connection_id is not None and str(connection_id).strip() != "":
            # Ensure the picker has connections loaded before selecting.
            if not self._ctx.connections():
                self._refresh_connections()
            selected = self._ctx.select_by_id(connection_id)
            if not selected:
                self._show_error(
                    f"Connection ID {connection_id} from session was not found in the connection list."
                )

    def _sync_target_table_for_active_sheet(self) -> None:
        if not self._sheets:
            return
        sheet = self._sheets[self._active_sheet_index]
        table_name = self._sheet_target_table_name(sheet if isinstance(sheet, dict) else None)
        if table_name:
            self.target_table_edit.setText(table_name)

    def _clear_analyze_panel(self) -> None:
        self._sheets = []
        self._saved_sheet_keys.clear()
        self._edit_mode = False
        self.edit_btn.setText("Edit mapping")
        self._clear_sheet_tabs()
        self.analyze_table.setRowCount(0)
        self.analyze_panel.setVisible(False)

    def _apply_analyze_sheets(self, sheets: list[Any], *, success_message: str | None = None) -> None:
        raw = sheets if isinstance(sheets, list) else []
        enriched = _enrich_analyze_sheets_unicode(
            raw,
            file_path=self._selected_path,
            multi_language=self.multi_language_enabled(),
        )
        self._sheets = deepcopy(enriched)
        self._saved_sheet_keys.clear()
        self._edit_mode = False
        self.edit_btn.setText("Edit mapping")
        if not self._sheets:
            self.analyze_panel.setVisible(False)
            return
        self._apply_sql_datatypes_for_connection()
        self._rebuild_sheet_tabs()
        self._set_active_sheet(0)
        self.analyze_panel.setVisible(True)
        if success_message:
            self._show_success(success_message)

    def _check_analyze_if_present(self, session_id: str) -> None:
        if self._uploading or self._analyzing:
            return
        self._analyzing = True
        self.upload_btn.setEnabled(False)
        try:
            result = api_analyze_import_if_present(session_id, token=self._token())
        finally:
            self._analyzing = False
            self.upload_btn.setEnabled(True)

        if not result.get("success"):
            # Soft-fail: keep normal upload flow available.
            self._clear_analyze_panel()
            self._clear_session_target_fields()
            self._show_error(str(result.get("message") or "Could not check session analysis."))
            return

        if result.get("present"):
            sheets = result.get("data") if isinstance(result.get("data"), list) else []
            sheets = [
                s
                for s in sheets
                if isinstance(s, dict)
                and isinstance(s.get("columns"), list)
                and s.get("columns")
            ]
            if sheets:
                self._clear_message()
                self._apply_session_target_meta(result, sheets)
                self._apply_analyze_sheets(
                    sheets,
                    success_message="Existing upload found. Workbook analysis loaded.",
                )
                return
        # Not uploaded yet — continue normal upload procedure.
        self._clear_analyze_panel()
        self._clear_session_target_fields()
        self._show_success(
            "No existing upload for this session. Select a file to upload."
        )

    def _upload_status_success(self, result: dict[str, Any]) -> bool:
        if not result.get("success"):
            return False
        data = result.get("data")
        if isinstance(data, dict):
            status = str(data.get("status") or result.get("status") or "").strip().upper()
            if status and status != "SUCCESS":
                return False
            return True
        status = str(result.get("status") or "").strip().upper()
        return (not status) or status == "SUCCESS"

    def selected_codepage(self) -> str:
        data = self.codepage_combo.currentData()
        if data is not None and str(data).strip():
            return str(data).strip()
        return "65001"

    def multi_language_enabled(self) -> bool:
        return bool(self.multi_language_check.isChecked())

    def _handle_upload(self) -> None:
        if self._uploading or self._analyzing:
            return
        session_id = self.selected_session_id()
        if not session_id:
            self._show_error("Select a session before uploading.")
            return
        if not self._selected_path:
            self._show_error("Select a file before uploading.")
            return

        self._clear_message()
        self._uploading = True
        self.upload_btn.setEnabled(False)
        try:
            result = api_upload_import_file(
                session_id,
                self._selected_path,
                token=self._token(),
                codepage=self.selected_codepage() if self.multi_language_enabled() else None,
                multi_language=self.multi_language_enabled(),
            )
        finally:
            self._uploading = False
            self.upload_btn.setEnabled(True)

        if not self._upload_status_success(result):
            self._show_error(str(result.get("message") or "Upload failed."))
            return

        self._show_success(str(result.get("message") or "File uploaded successfully."))
        self._run_analyze(session_id)

    def _run_analyze(self, session_id: str) -> None:
        self._analyzing = True
        self.upload_btn.setEnabled(False)
        try:
            result = api_analyze_import_file(
                session_id,
                token=self._token(),
                codepage=self.selected_codepage() if self.multi_language_enabled() else None,
                multi_language=self.multi_language_enabled(),
            )
        finally:
            self._analyzing = False
            self.upload_btn.setEnabled(True)

        if not result.get("success"):
            self._show_error(str(result.get("message") or "Analyze failed."))
            self.analyze_panel.setVisible(False)
            return

        sheets = result.get("data") if isinstance(result.get("data"), list) else []
        if not sheets:
            self.analyze_panel.setVisible(False)
            self._show_error("Analyze succeeded but returned no sheets.")
            return
        self._apply_analyze_sheets(sheets)

    def _clear_sheet_tabs(self) -> None:
        while self.sheet_tabs_layout.count():
            item = self.sheet_tabs_layout.takeAt(0)
            w = item.widget()
            if w is not None:
                w.deleteLater()
        self._sheet_tab_buttons = []

    def _rebuild_sheet_tabs(self) -> None:
        self._clear_sheet_tabs()
        for index, sheet in enumerate(self._sheets):
            name = str(sheet.get("name") or f"Sheet{index + 1}")
            btn = QPushButton(name)
            btn.setCheckable(True)
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.setStyleSheet(self._sheet_tab_style_for(sheet, index))
            btn.clicked.connect(lambda checked=False, i=index: self._set_active_sheet(i))
            self.sheet_tabs_layout.addWidget(btn)
            self._sheet_tab_buttons.append(btn)
        self.sheet_tabs_layout.addStretch()

    def _sheet_tab_style_for(self, sheet: dict[str, Any], index: int) -> str:
        key = self._sheet_key(sheet, index)
        return _SHEET_TAB_SAVED_STYLE if key in self._saved_sheet_keys else _SHEET_TAB_STYLE

    def _refresh_sheet_tab_styles(self) -> None:
        for index, btn in enumerate(self._sheet_tab_buttons):
            sheet = self._sheets[index] if index < len(self._sheets) else {}
            btn.setStyleSheet(self._sheet_tab_style_for(sheet, index))
        if 0 <= self._active_sheet_index < len(self._sheets):
            sheet = self._sheets[self._active_sheet_index]
            sheet_name = str(sheet.get("name") or f"Sheet{self._active_sheet_index + 1}")
            self.choose_sheet_chip.setText(sheet_name)
            if self._sheet_key(sheet, self._active_sheet_index) in self._saved_sheet_keys:
                self.choose_sheet_chip.setStyleSheet(
                    "QLabel#chooseSheetChip { background: #dcfce7; color: #15803d; "
                    "border: 1px solid #86efac; border-radius: 10px; padding: 3px 10px; "
                    "font-size: 11px; font-weight: 700; }"
                )
            else:
                self.choose_sheet_chip.setStyleSheet(_CHOOSE_SHEET_STYLE)

    def _set_active_sheet(self, index: int) -> None:
        if index < 0 or index >= len(self._sheets):
            return
        if self._edit_mode:
            self._capture_table_edits_into_sheet()
        self._active_sheet_index = index
        for i, btn in enumerate(self._sheet_tab_buttons):
            btn.setChecked(i == index)
        self._refresh_sheet_tab_styles()
        self._sync_target_table_for_active_sheet()
        self._render_active_sheet()

    def _render_active_sheet(self) -> None:
        sheet = self._sheets[self._active_sheet_index] if self._sheets else {}
        columns = sheet.get("columns") if isinstance(sheet, dict) else []
        if not isinstance(columns, list):
            columns = []

        self.analyze_table.setRowCount(0)
        self.analyze_table.setRowCount(len(columns))
        self.analyze_table.setVerticalHeaderLabels(
            [str(row_number) for row_number in range(1, len(columns) + 1)]
        )
        for row_index, col in enumerate(columns):
            if not isinstance(col, dict):
                col = {}
            db_type = self._selected_db_type()
            datatype_options = self._datatype_options()
            for col_index, (key, _label) in enumerate(_ANALYZE_COLUMNS):
                value = col.get(key, "")
                text = "" if value is None else str(value)
                if key == "tgtDatatype":
                    sql_text = import_target_data_type_to_sql(text or col.get("datatype") or "", db_type)
                    if self._edit_mode:
                        combo = QComboBox()
                        for opt in datatype_options:
                            combo.addItem(opt)
                        if sql_text and combo.findText(sql_text) < 0:
                            combo.insertItem(0, sql_text)
                        apply_form_combobox_field(
                            combo,
                            height_px=FORM_SINGLELINE_FIELD_HEIGHT_PX,
                        )
                        install_combo_popup_below_field(combo)
                        idx = combo.findText(sql_text)
                        combo.setCurrentIndex(idx if idx >= 0 else 0)
                        self.analyze_table.setCellWidget(row_index, col_index, combo)
                        continue
                    item = QTableWidgetItem(sql_text)
                    item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                    item.setBackground(QColor("#f8fbff"))
                    self.analyze_table.setItem(row_index, col_index, item)
                    continue

                item = QTableWidgetItem(text)
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                if self._edit_mode and key in ("tgtColumnName", "length", "decimal", "precision"):
                    item.setFlags(item.flags() | Qt.ItemFlag.ItemIsEditable)
                    item.setBackground(QColor("#eff6ff"))
                elif key in ("tgtColumnName", "length", "decimal", "precision"):
                    item.setBackground(QColor("#f8fbff"))
                if key in ("ordinal", "length", "decimal", "precision"):
                    item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                name_text = str(col.get("name") or "")
                if key in ("name", "tgtColumnName") and name_text.startswith("#"):
                    item.setForeground(QColor("#dc2626"))
                self.analyze_table.setItem(row_index, col_index, item)

        if self._edit_mode:
            self.analyze_table.setEditTriggers(
                QAbstractItemView.EditTrigger.DoubleClicked
                | QAbstractItemView.EditTrigger.EditKeyPressed
                | QAbstractItemView.EditTrigger.AnyKeyPressed
            )
        else:
            self.analyze_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        for column_index, width in enumerate((150, 80, 190, 140, 80, 80, 90)):
            self.analyze_table.setColumnWidth(column_index, width)

    def _capture_table_edits_into_sheet(self) -> None:
        if not self._sheets:
            return
        sheet = self._sheets[self._active_sheet_index]
        columns = sheet.get("columns")
        if not isinstance(columns, list):
            return
        for row_index, col in enumerate(columns):
            if not isinstance(col, dict):
                continue
            for col_index, (key, _label) in enumerate(_ANALYZE_COLUMNS):
                if key == "tgtDatatype":
                    widget = self.analyze_table.cellWidget(row_index, col_index)
                    if isinstance(widget, QComboBox):
                        col[key] = import_target_data_type_to_sql(
                            widget.currentText(),
                            self._selected_db_type(),
                        )
                    continue
                if key not in ("tgtColumnName", "length", "decimal", "precision"):
                    continue
                item = self.analyze_table.item(row_index, col_index)
                if item is not None:
                    col[key] = item.text().strip()

    def _toggle_edit_mode(self) -> None:
        if not self.analyze_panel.isVisible() or not self._sheets:
            return
        if self._edit_mode:
            self._capture_table_edits_into_sheet()
            self._edit_mode = False
            self.edit_btn.setText("Edit mapping")
        else:
            self._edit_mode = True
            self.edit_btn.setText("Done editing")
        self._render_active_sheet()

    def _resolve_sheet_id(self, sheet: dict[str, Any]) -> int | None:
        raw = sheet.get("_raw") if isinstance(sheet.get("_raw"), dict) else {}
        value = sheet.get("sheetId")
        if value is None:
            value = _first_mapping_id(raw, "sheetId", "sheet_id", "sheetID", "id")
        try:
            return int(value) if value is not None and str(value).strip() != "" else None
        except (TypeError, ValueError):
            return None

    def _resolve_source_column_id(self, col: dict[str, Any]) -> int | None:
        raw = col.get("_raw") if isinstance(col.get("_raw"), dict) else {}
        value = col.get("sourceColumnId")
        if value is None:
            value = _first_mapping_id(
                raw,
                "sourceColumnId",
                "source_column_id",
                "columnId",
                "column_id",
                "fieldId",
                "id",
            )
        try:
            return int(value) if value is not None and str(value).strip() != "" else None
        except (TypeError, ValueError):
            return None

    def _build_mapping_body(self) -> tuple[dict[str, Any] | None, str | None]:
        if self._edit_mode:
            self._capture_table_edits_into_sheet()

        target_table = self.selected_target_table_name()
        if not target_table:
            return None, "Enter a target table name before saving."

        conn_id = etl_connection_id(self.selected_connection())
        if conn_id is None or str(conn_id).strip() == "":
            return None, "Select a connection before saving."
        try:
            target_connection_id = int(conn_id)
        except (TypeError, ValueError):
            return None, "Connection ID must be numeric."

        if not self._sheets:
            return None, "Analyze a workbook before saving the mapping."
        sheet = self._sheets[self._active_sheet_index]
        sheet_id = self._resolve_sheet_id(sheet)
        if sheet_id is None:
            return None, "Sheet ID is missing from analyze result."

        columns = sheet.get("columns") if isinstance(sheet, dict) else []
        if not isinstance(columns, list) or not columns:
            return None, "No columns available to save."

        mappings: list[dict[str, Any]] = []
        for col in columns:
            if not isinstance(col, dict):
                continue
            source_column_id = self._resolve_source_column_id(col)
            if source_column_id is None:
                return None, "One or more columns are missing sourceColumnId."
            target_column_name = str(col.get("tgtColumnName") or "").strip()
            if not target_column_name:
                return None, "Every mapping needs a target column name."
            java_type = import_target_data_type_to_java(col.get("tgtDatatype") or col.get("datatype"))
            nullable = bool(col.get("nullableBool", True))
            if col.get("nullable") in ("FALSE", "false", "0", False):
                nullable = False
            elif col.get("nullable") in ("TRUE", "true", "1", True):
                nullable = True
            mapping: dict[str, Any] = {
                "sourceColumnId": source_column_id,
                "targetColumnName": target_column_name,
                "targetDataType": java_type,
                "nullable": nullable,
                "primaryKey": bool(col.get("primaryKey", False)),
                "selected": bool(col.get("selected", True)),
            }
            length_raw = str(col.get("length") or "").strip()
            if length_raw and java_type in ("STRING", "DECIMAL"):
                try:
                    mapping["length"] = int(length_raw)
                except ValueError:
                    pass
            mappings.append(mapping)

        if not mappings:
            return None, "No mappings available to save."

        return (
            {
                "sheetId": sheet_id,
                "targetTableName": target_table,
                "targetConnectionId": target_connection_id,
                "mappings": mappings,
            },
            None,
        )

    def _handle_save(self) -> None:
        if self._saving_mapping or self._uploading or self._analyzing:
            return
        session_id = self.selected_session_id()
        if not session_id:
            self._show_error("Select a session before saving.")
            return
        body, error = self._build_mapping_body()
        if error or body is None:
            self._show_error(error or "Could not build mapping payload.")
            return

        if self._edit_mode:
            self._edit_mode = False
            self.edit_btn.setText("Edit mapping")
            self._render_active_sheet()

        self._clear_message()
        self._saving_mapping = True
        self.footer_save_btn.setEnabled(False)
        self.save_btn.setEnabled(False)
        try:
            result = api_save_import_mapping(session_id, body, token=self._token())
        finally:
            self._saving_mapping = False
            self.footer_save_btn.setEnabled(True)
            self.save_btn.setEnabled(True)

        if not result.get("success"):
            self._show_error(str(result.get("message") or "Failed to save mapping."))
            return
        sheet = self._sheets[self._active_sheet_index]
        self._saved_sheet_keys.add(self._sheet_key(sheet, self._active_sheet_index))
        self._refresh_sheet_tab_styles()
        self._show_success(str(result.get("message") or "Mapping saved successfully."))


def _first_mapping_id(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and row.get(key) is not None:
            return row.get(key)
    return None
