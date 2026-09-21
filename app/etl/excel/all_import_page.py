"""ETL: Excel — All Import (Upload & Analyze / Map Columns / Execute)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from PySide6.QtCore import Qt, QTimer, Signal
from PySide6.QtGui import QDragEnterEvent, QDropEvent, QPainter, QPixmap
from PySide6.QtSvg import QSvgRenderer
from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QProgressBar,
    QPushButton,
    QRadioButton,
    QSizePolicy,
    QStackedWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.etl.excel.message_banner import (
    clear_page_message,
    configure_message_banner,
    show_page_error,
    show_page_info,
    show_page_success,
)
from core.api import (
    api_analyze_import_file,
    api_analyze_import_if_present,
    api_create_import,
    api_execute_import_sheet,
    api_get_import_sheets_by_uuid,
    api_save_import_mapping,
    api_upload_import_file,
    api_validate_import_sheet,
    import_target_data_type_to_java,
)
from core.etl_connection_context import (
    etl_connection_display_label,
    etl_connection_id,
    get_etl_connection_context,
)
from core.user_context import get_user_profile
from ui.data_table import apply_data_table_appearance, attach_table_copy_shortcut
from ui.form_combobox_style import apply_form_combobox_field, install_combo_popup_below_field
from ui.form_page_styles import FORM_INPUT_STYLE
from ui.theme import Theme

_ASSETS = Path(__file__).resolve().parents[3] / "assets" / "excel"
_UPLOAD_CLOUD_SVG = _ASSETS / "upload_cloud.svg"
_FILE_SPREADSHEET_SVG = _ASSETS / "file_spreadsheet.svg"
_CHECK_GREEN_SVG = _ASSETS / "check_green.svg"
_FILE_WHITE_SVG = _ASSETS / "file_white.svg"
_KEY_BLUE_SVG = _ASSETS / "key_blue.svg"
_LOADER_SVG = _ASSETS / "loader.svg"
_CHECK_SUCCESS_SVG = _ASSETS / "check_success.svg"

_FILE_FILTER = "Spreadsheets (*.xlsx *.xls *.csv);;Excel (*.xlsx *.xls);;CSV (*.csv);;All files (*.*)"
_ALLOWED_SUFFIXES = {".xlsx", ".xls", ".csv"}

_SOURCE_TYPE_OPTIONS: tuple[tuple[str, str], ...] = (
    ("Excel", "EXCEL"),
    ("CSV", "CSV"),
    ("JSON", "JSON"),
)

_ENCODING_OPTIONS: tuple[tuple[str, str], ...] = (
    ("UTF-8", "UTF_8"),
    ("Windows Arabic — 1256", "WINDOWS_1256"),
    ("Windows Western European — 1252", "WINDOWS_1252"),
)

_DATATYPE_OPTIONS = ("INT", "BIGINT", "NVARCHAR", "VARCHAR", "DECIMAL", "FLOAT", "BIT", "DATE", "DATETIME", "TEXT")

_PAGE_BG = "#f7f7fa"
_ACCENT = "#2e73f5"
_TEXT_DARK = "#111827"
_TEXT_MUTED = "#6b7280"
_BORDER = "#e8eaed"
_BORDER_INPUT = "#d1d6de"
_CARD_BG = "#f6f7f9"

_PAGE_STYLE = f"""
    QWidget#allImportBody {{ background: {_PAGE_BG}; }}
    QLabel#allImportTitle {{
        color: {_TEXT_DARK}; font-size: 22px; font-weight: 700;
        background: transparent; border: none;
    }}
    QPushButton#allImportPrimary {{
        background: {_ACCENT}; color: #ffffff; border: none; border-radius: 4px;
        padding: 10px 24px; font-size: 13px; font-weight: 600;
    }}
    QPushButton#allImportPrimary:hover {{ background: #1f63e0; }}
    QPushButton#allImportSecondary {{
        background: #ffffff; color: {_TEXT_MUTED}; border: 1px solid {_BORDER_INPUT};
        border-radius: 4px; padding: 10px 24px; font-size: 13px; font-weight: 600;
    }}
    QPushButton#allImportSecondary:hover {{ background: #f8fafc; border-color: #94a3b8; }}
    QFrame#allImportTabBar {{
        background: transparent; border: none; border-bottom: 2px solid {_BORDER};
    }}
    QWidget#allImportTab {{
        background: transparent; border: none;
    }}
    QPushButton#allImportTabHit {{
        background: transparent;
        border: none;
        border-bottom: 2px solid transparent;
        padding: 8px 18px;
        margin-bottom: -2px;
        font-size: 13px;
        font-weight: 500;
        color: {_TEXT_MUTED};
    }}
    QPushButton#allImportTabHit:checked {{
        color: {_ACCENT};
        font-weight: 600;
        border-bottom: 2px solid {_ACCENT};
    }}
    QPushButton#allImportTabHit:hover {{
        color: {_ACCENT};
    }}
    QFrame#allImportLeftCard, QFrame#allImportRightCard {{
        background: {_CARD_BG}; border: 1px solid {_BORDER}; border-radius: 8px;
    }}
    QLabel#allImportSectionTitle {{
        color: {_TEXT_DARK}; font-size: 14px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportFieldCaption {{
        color: {_TEXT_MUTED}; font-size: 11px; font-weight: 600;
        background: transparent; border: none;
    }}
    QFrame#allImportDropZone {{
        background: #ffffff; border: 1px dashed {_TEXT_MUTED}; border-radius: 8px;
    }}
    QFrame#allImportDropZone[dragActive="true"] {{
        border: 1px dashed {_ACCENT}; background: #eff6ff;
    }}
    QLabel#allImportDropTitle {{
        color: {_TEXT_DARK}; font-size: 14px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportDropHint {{
        color: {_TEXT_MUTED}; font-size: 12px; background: transparent; border: none;
    }}
    QFrame#allImportUploadedBanner {{
        background: #ffffff; border: 1px solid {_BORDER}; border-radius: 6px;
    }}
    QLabel#allImportUploadedName {{
        color: {_TEXT_DARK}; font-size: 14px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportUploadedSize {{
        color: {_TEXT_MUTED}; font-size: 12px; background: transparent; border: none;
    }}
    QFrame#allImportSuccessPill {{
        background: #dcfce7; border: none; border-radius: 100px;
    }}
    QLabel#allImportSuccessPillText {{
        color: #22c55e; font-size: 11px; font-weight: 600;
        background: transparent; border: none;
    }}
    QPushButton#allImportChangeFile {{
        background: #ffffff; color: {_ACCENT}; border: 1px solid {_BORDER_INPUT};
        border-radius: 4px; padding: 6px 12px; font-size: 12px; font-weight: 600;
    }}
    QPushButton#allImportChangeFile:hover {{
        background: #eff6ff; border-color: {_ACCENT};
    }}
    QLabel#allImportAnalysisHeading {{
        color: {_TEXT_MUTED}; font-size: 12px; font-weight: 700;
        background: transparent; border: none;
    }}
    QFrame#allImportAnalysisBadge {{
        background: #dcfce7; border: none; border-radius: 4px;
    }}
    QLabel#allImportAnalysisBadgeText {{
        color: #22c55e; font-size: 11px; font-weight: 600;
        background: transparent; border: none;
    }}
    QFrame#allImportAnalysisCard {{
        background: #ffffff; border: 1px solid {_BORDER}; border-radius: 6px;
    }}
    QLabel#allImportAnalysisLine {{
        color: {_TEXT_DARK}; font-size: 13px; background: transparent; border: none;
    }}
    QFrame#allImportFileSummary {{
        background: #1c212b; border: none; border-radius: 6px;
    }}
    QLabel#allImportFileSummaryText {{
        color: #ffffff; font-size: 13px; font-weight: 500;
        background: transparent; border: none;
    }}
    QFrame#allImportModeGroup {{
        background: #ffffff; border: 1px solid {_BORDER_INPUT}; border-radius: 6px;
    }}
    QPushButton#allImportModeBtn {{
        background: transparent; border: none; border-radius: 4px;
        padding: 6px 16px; font-size: 12px; font-weight: 500; color: {_TEXT_MUTED};
    }}
    QPushButton#allImportModeBtn:checked {{
        background: {_ACCENT}; color: #ffffff; font-weight: 600;
    }}
    QPushButton#allImportSheetBtn {{
        background: #ffffff; border: 1px solid {_BORDER_INPUT}; border-radius: 4px;
        padding: 6px 12px; font-size: 12px; font-weight: 500; color: {_TEXT_MUTED};
    }}
    QPushButton#allImportSheetBtn:checked {{
        background: {_ACCENT}; color: #ffffff; border-color: {_ACCENT}; font-weight: 600;
    }}
    QPushButton#allImportPrimary:disabled {{
        background: #94a1b0; color: #ffffff;
    }}
    QFrame#allImportProgressCard {{
        background: #ffffff; border: 1px solid {_BORDER}; border-radius: 6px;
    }}
    QLabel#allImportProgressStatus {{
        color: {_TEXT_DARK}; font-size: 14px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportProgressPct {{
        color: {_ACCENT}; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QProgressBar#allImportProgressBar {{
        background: {_CARD_BG}; border: none; border-radius: 4px; min-height: 8px; max-height: 8px;
        text-align: center;
    }}
    QProgressBar#allImportProgressBar::chunk {{
        background: {_ACCENT}; border-radius: 4px;
    }}
    QLabel#allImportMetricLabel {{
        color: {_TEXT_MUTED}; font-size: 13px; font-weight: 400;
        background: transparent; border: none;
    }}
    QLabel#allImportMetricValue {{
        color: {_TEXT_DARK}; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportMetricSuccess {{
        color: #22c55e; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportMetricFailed {{
        color: #d92e2e; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QRadioButton#allImportOpRadio {{
        color: {_TEXT_MUTED}; font-size: 13px; font-weight: 400; spacing: 8px;
    }}
    QRadioButton#allImportOpRadio:checked {{
        color: {_TEXT_DARK}; font-weight: 600;
    }}
    QLineEdit#allImportReadonlyField {{
        background: #ffffff; border: 1px solid {_BORDER_INPUT}; border-radius: 4px;
        padding: 0 12px; font-size: 13px; color: {_TEXT_DARK};
    }}
    QFrame#allImportSuccessSummary {{
        background: {_CARD_BG}; border: 1px solid {_BORDER}; border-radius: 8px;
    }}
    QLabel#allImportSuccessTitle {{
        color: {_TEXT_DARK}; font-size: 20px; font-weight: 700;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessSubtitle {{
        color: {_TEXT_MUTED}; font-size: 14px; font-weight: 400;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessSummaryHeading {{
        color: {_TEXT_MUTED}; font-size: 13px; font-weight: 700;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessRowLabel {{
        color: {_TEXT_MUTED}; font-size: 13px; font-weight: 400;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessRowValue {{
        color: {_TEXT_DARK}; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessRowSuccess {{
        color: #22c55e; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportSuccessRowFailed {{
        color: #d92e2e; font-size: 13px; font-weight: 600;
        background: transparent; border: none;
    }}
    QLabel#allImportStatusBadge {{
        background: #dcfce7; color: #22c55e; font-size: 11px; font-weight: 600;
        border: none; border-radius: 4px; padding: 2px 8px;
    }}
    QPushButton#allImportViewDetails {{
        background: #ffffff; color: {_TEXT_MUTED}; border: 1px solid {_BORDER_INPUT};
        border-radius: 4px; padding: 0 20px; font-size: 13px; font-weight: 600;
        min-height: 36px; max-height: 36px;
    }}
    QPushButton#allImportViewDetails:hover {{
        background: {_CARD_BG};
    }}
    QPushButton#allImportBackToImports {{
        background: {_ACCENT}; color: #ffffff; border: none;
        border-radius: 4px; padding: 0 20px; font-size: 13px; font-weight: 600;
        min-height: 36px; max-height: 36px;
    }}
    QPushButton#allImportBackToImports:hover {{
        background: #2563eb;
    }}
    QPushButton#allImportNewImport {{
        background: #ffffff; color: {_ACCENT}; border: 1px solid {_ACCENT};
        border-radius: 4px; padding: 0 20px; font-size: 13px; font-weight: 600;
        min-height: 36px; max-height: 36px;
    }}
    QPushButton#allImportNewImport:hover {{
        background: {_CARD_BG};
    }}
    QTableWidget#allImportMapTable {{
        background: #ffffff; alternate-background-color: {_CARD_BG};
        border: 1px solid {_BORDER}; border-radius: 8px; gridline-color: {_BORDER};
        font-size: 13px; color: {_TEXT_DARK}; outline: none;
    }}
    QTableWidget#allImportMapTable::item {{
        padding: 4px 8px; border: none;
    }}
    QHeaderView#allImportMapHeader::section {{
        background: {_CARD_BG}; color: {_TEXT_MUTED}; font-size: 11px; font-weight: 700;
        border: none; border-bottom: 1px solid {_BORDER}; border-right: 1px solid {_BORDER};
        padding: 0 8px;
    }}
    QFrame#allImportSuccessDivider {{
        background: {_BORDER}; border: none; max-height: 1px; min-height: 1px;
    }}
"""


def _svg_pixmap(path: Path, size: int) -> QPixmap:
    pm = QPixmap(size, size)
    pm.fill(Qt.GlobalColor.transparent)
    if path.is_file():
        renderer = QSvgRenderer(str(path))
        painter = QPainter(pm)
        renderer.render(painter)
        painter.end()
    return pm


def _format_file_size(num_bytes: int) -> str:
    if num_bytes < 1024:
        return f"{num_bytes} B"
    kb = num_bytes / 1024
    if kb < 1024:
        return f"{kb:.1f} KB"
    return f"{kb / 1024:.1f} MB"


def _to_snake_case(name: str) -> str:
    text = re.sub(r"[^0-9A-Za-z]+", "_", name.strip())
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", text)
    text = re.sub(r"_+", "_", text).strip("_").lower()
    return text or "column"


def _bracket_spaced_identifier(name: str) -> str:
    """Wrap identifiers that contain spaces in SQL-style ``[brackets]``.

    Needed so target column names with spaces succeed on create/execute.
    """
    text = str(name or "").strip()
    if not text:
        return text
    if text.startswith("[") and text.endswith("]") and len(text) >= 2:
        return text
    if " " not in text:
        return text
    return f"[{text.replace(']', ']]')}]"


def _map_centered_widget(child: QWidget) -> QWidget:
    wrap = QWidget()
    wrap.setStyleSheet("background: transparent; border: none;")
    lay = QHBoxLayout(wrap)
    lay.setContentsMargins(0, 0, 0, 0)
    lay.setAlignment(Qt.AlignmentFlag.AlignCenter)
    lay.addWidget(child)
    return wrap


def _guess_datatype(header: str) -> tuple[str, str]:
    lower = header.lower()
    if any(token in lower for token in ("id", "qty", "quantity", "count", "total", "price", "amount")):
        return "INT", "-"
    if any(token in lower for token in ("date", "time")):
        return "DATETIME", "-"
    if any(token in lower for token in ("name", "product", "category", "size", "desc")):
        return "NVARCHAR", "255"
    return "NVARCHAR", "100"


def _session_id_from_create(result: dict[str, Any]) -> str:
    data = result.get("data")
    candidates: list[Any] = []
    if isinstance(data, dict):
        candidates.append(data)
        nested = data.get("data")
        if isinstance(nested, dict):
            candidates.append(nested)
    candidates.append(result)
    keys = (
        "sessionId",
        "sessionID",
        "session_id",
        "importSessionId",
        "importId",
        "uuid",
        "id",
    )
    for obj in candidates:
        if not isinstance(obj, dict):
            continue
        for key in keys:
            value = obj.get(key)
            if value is not None and str(value).strip():
                return str(value).strip()
    return ""


def _first_id(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and row.get(key) is not None:
            return row.get(key)
    raw = row.get("_raw")
    if isinstance(raw, dict):
        for key in keys:
            if key in raw and raw.get(key) is not None:
                return raw.get(key)
    return None


def _encoding_codepage(encoding_api: str) -> str | None:
    """Map API encoding enum to optional multipart codepage hint."""
    key = str(encoding_api or "").strip().upper().replace("-", "_")
    mapping = {
        "UTF_8": "65001",
        "UTF8": "65001",
        "WINDOWS_1256": "1256",
        "WINDOWS_1252": "1252",
        "1256": "1256",
        "1252": "1252",
        "65001": "65001",
    }
    return mapping.get(key)


def _analyze_workbook(path: str) -> list[dict[str, object]]:
    """Local preview only ``[{name, headers: list[str]}, ...]`` before API analyze."""
    suffix = Path(path).suffix.lower()
    if suffix == ".csv":
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                first = handle.readline()
            headers = [h.strip().strip('"') for h in first.split(",") if h.strip()]
            return [{"name": "Sheet1", "headers": headers}]
        except OSError:
            return [{"name": "Sheet1", "headers": []}]
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    sheets: list[dict[str, object]] = []
    try:
        for ws in wb.worksheets:
            headers: list[str] = []
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    text = str(cell).strip()
                    if text:
                        headers.append(text)
                break
            sheets.append({"name": str(ws.title or "Sheet"), "headers": headers})
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return sheets


class _FileDropZone(QFrame):
    file_selected = Signal(str)

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setObjectName("allImportDropZone")
        self.setAcceptDrops(True)
        self.setCursor(Qt.CursorShape.PointingHandCursor)
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Expanding)
        self.setProperty("dragActive", False)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(24, 48, 24, 48)
        layout.setSpacing(12)
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon = QLabel()
        icon.setFixedSize(32, 32)
        icon.setPixmap(_svg_pixmap(_UPLOAD_CLOUD_SVG, 32))
        layout.addWidget(icon, 0, Qt.AlignmentFlag.AlignHCenter)
        title = QLabel("Drag & drop your file here or click to browse")
        title.setObjectName("allImportDropTitle")
        title.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        title.setWordWrap(True)
        layout.addWidget(title)
        hint = QLabel("Supported formats: .xlsx, .xls, .csv (Max 10MB)")
        hint.setObjectName("allImportDropHint")
        hint.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        hint.setWordWrap(True)
        layout.addWidget(hint)

    def mousePressEvent(self, event) -> None:  # noqa: N802
        if event.button() == Qt.MouseButton.LeftButton:
            self._browse()
        super().mousePressEvent(event)

    def dragEnterEvent(self, event: QDragEnterEvent) -> None:  # noqa: N802
        if event.mimeData().hasUrls():
            for url in event.mimeData().urls():
                if Path(url.toLocalFile()).suffix.lower() in _ALLOWED_SUFFIXES:
                    event.acceptProposedAction()
                    self._set_drag_active(True)
                    return
        event.ignore()

    def dragLeaveEvent(self, event) -> None:  # noqa: N802
        self._set_drag_active(False)
        super().dragLeaveEvent(event)

    def dropEvent(self, event: QDropEvent) -> None:  # noqa: N802
        self._set_drag_active(False)
        for url in event.mimeData().urls():
            path = url.toLocalFile()
            if Path(path).suffix.lower() in _ALLOWED_SUFFIXES:
                self.file_selected.emit(path)
                event.acceptProposedAction()
                return
        event.ignore()

    def _set_drag_active(self, active: bool) -> None:
        self.setProperty("dragActive", active)
        self.style().unpolish(self)
        self.style().polish(self)

    def _browse(self) -> None:
        path, _ = QFileDialog.getOpenFileName(self, "Select file", "", _FILE_FILTER)
        if path:
            self.file_selected.emit(path)


class AllImportPage(QWidget):
    """Excel All Import wizard: Upload & Analyze → Map Columns → Execute."""

    exit_to_list_requested = Signal()

    def __init__(self, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self._selected_path = ""
        self._file_size_text = ""
        self._step_index = 0
        self._session_id = ""
        self._sheets: list[dict[str, Any]] = []
        self._preview_sheets: list[dict[str, object]] = []
        self._active_sheet = 0
        self._sheet_buttons: list[QPushButton] = []
        self._busy = False
        self._importing = False
        self._import_complete = False
        self._remote_session = False
        self._progress_value = 0
        self._total_rows = 0
        self._progress_timer = QTimer(self)
        self._progress_timer.setInterval(120)
        self._progress_timer.timeout.connect(self._tick_import_progress)
        self._ctx = get_etl_connection_context()
        self._ctx.connections_changed.connect(self._refresh_connections)

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        body = QWidget()
        body.setObjectName("allImportBody")
        body.setStyleSheet(_PAGE_STYLE)
        body_l = QVBoxLayout(body)
        body_l.setContentsMargins(46, 20, 46, 24)
        body_l.setSpacing(16)

        # Header: title left, actions right
        header_row = QHBoxLayout()
        header_row.setContentsMargins(0, 0, 0, 0)
        header_row.setSpacing(12)
        header_row.setAlignment(Qt.AlignmentFlag.AlignVCenter)
        title = QLabel("Excel Import")
        title.setObjectName("allImportTitle")
        header_row.addWidget(title, 1, Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        btn_row = QHBoxLayout()
        btn_row.setContentsMargins(0, 0, 0, 0)
        btn_row.setSpacing(12)
        self.primary_btn = QPushButton("Import")
        self.primary_btn.setObjectName("allImportPrimary")
        self.primary_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.primary_btn.clicked.connect(self._on_primary)
        btn_row.addWidget(self.primary_btn)
        self.secondary_btn = QPushButton("Cancel")
        self.secondary_btn.setObjectName("allImportSecondary")
        self.secondary_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.secondary_btn.clicked.connect(self._on_secondary)
        btn_row.addWidget(self.secondary_btn)
        header_row.addLayout(btn_row)
        body_l.addLayout(header_row)

        # Step tabs — directly under the Excel Import heading
        tab_bar = QFrame()
        tab_bar.setObjectName("allImportTabBar")
        tab_bar.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
        tab_l = QHBoxLayout(tab_bar)
        tab_l.setContentsMargins(0, 0, 0, 0)
        tab_l.setSpacing(0)
        tab_l.setAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignBottom)
        self._tab_buttons: list[QPushButton] = []
        self._tab_step_labels = ("Upload & Analyze", "Map Columns", "Execute & Result")
        for index, label in enumerate(self._tab_step_labels):
            tab = QPushButton(f"  {index + 1}   {label}")
            tab.setObjectName("allImportTabHit")
            tab.setCheckable(True)
            tab.setChecked(index == 0)
            tab.setCursor(Qt.CursorShape.PointingHandCursor)
            tab.setFlat(True)
            tab.setMinimumHeight(36)
            tab.clicked.connect(lambda _c=False, i=index: self._on_tab_clicked(i))
            tab_l.addWidget(tab)
            self._tab_buttons.append(tab)
        tab_l.addStretch(1)
        body_l.addWidget(tab_bar)

        self.message_label = QLabel()
        self.message_label.setObjectName("allImportMessageBanner")
        configure_message_banner(self.message_label)
        body_l.addWidget(self.message_label)

        self.stack = QStackedWidget()
        self.stack.addWidget(self._build_upload_step())
        self.stack.addWidget(self._build_map_step())
        self.stack.addWidget(self._build_execute_step())
        body_l.addWidget(self.stack, 1)

        root.addWidget(body, 1)
        self._refresh_connections()
        self._apply_step_chrome()

    def _build_upload_step(self) -> QWidget:
        page = QWidget()
        split = QHBoxLayout(page)
        split.setContentsMargins(0, 0, 0, 0)
        split.setSpacing(24)

        left = QFrame()
        left.setObjectName("allImportLeftCard")
        left_l = QVBoxLayout(left)
        left_l.setContentsMargins(16, 16, 16, 16)
        left_l.setSpacing(16)
        left_title = QLabel("Import Configuration")
        left_title.setObjectName("allImportSectionTitle")
        left_l.addWidget(left_title)

        left_l.addWidget(self._field_caption("IMPORT NAME *"))
        self.import_name_edit = QLineEdit()
        self.import_name_edit.setPlaceholderText("Enter session name")
        self.import_name_edit.setStyleSheet(FORM_INPUT_STYLE)
        self.import_name_edit.setFixedHeight(40)
        left_l.addWidget(self.import_name_edit)

        type_enc = QHBoxLayout()
        type_enc.setSpacing(16)
        type_col = QVBoxLayout()
        type_col.setSpacing(6)
        type_col.addWidget(self._field_caption("FILE TYPE *"))
        self.file_type_combo = QComboBox()
        for label, value in _SOURCE_TYPE_OPTIONS:
            self.file_type_combo.addItem(label, value)
        self.file_type_combo.setCurrentIndex(0)
        apply_form_combobox_field(self.file_type_combo, height_px=40, min_width=120)
        install_combo_popup_below_field(self.file_type_combo)
        type_col.addWidget(self.file_type_combo)
        type_enc.addLayout(type_col, 1)

        enc_col = QVBoxLayout()
        enc_col.setSpacing(6)
        enc_col.addWidget(self._field_caption("ENCODING *"))
        self.encoding_combo = QComboBox()
        for label, value in _ENCODING_OPTIONS:
            self.encoding_combo.addItem(label, value)
        apply_form_combobox_field(self.encoding_combo, height_px=40, min_width=120)
        install_combo_popup_below_field(self.encoding_combo)
        enc_col.addWidget(self.encoding_combo)
        type_enc.addLayout(enc_col, 1)
        left_l.addLayout(type_enc)

        left_l.addWidget(self._field_caption("TARGET DATABASE CONNECTION *"))
        self.connection_combo = QComboBox()
        self.connection_combo.setPlaceholderText("Select connection")
        apply_form_combobox_field(self.connection_combo, height_px=40, min_width=200)
        install_combo_popup_below_field(self.connection_combo)
        left_l.addWidget(self.connection_combo)

        left_l.addWidget(self._field_caption("TARGET TABLE NAME *"))
        self.target_table_edit = QLineEdit()
        self.target_table_edit.setPlaceholderText("Enter target table name")
        self.target_table_edit.setStyleSheet(FORM_INPUT_STYLE)
        self.target_table_edit.setFixedHeight(40)
        left_l.addWidget(self.target_table_edit)
        left_l.addStretch(1)
        split.addWidget(left, 1)

        right = QFrame()
        right.setObjectName("allImportRightCard")
        right_l = QVBoxLayout(right)
        right_l.setContentsMargins(16, 16, 16, 16)
        right_l.setSpacing(16)
        right_title = QLabel("File Upload")
        right_title.setObjectName("allImportSectionTitle")
        right_l.addWidget(right_title)

        self.drop_zone = _FileDropZone()
        self.drop_zone.file_selected.connect(self._on_file_selected)
        right_l.addWidget(self.drop_zone, 1)

        self._uploaded_panel = QWidget()
        uploaded_l = QVBoxLayout(self._uploaded_panel)
        uploaded_l.setContentsMargins(0, 0, 0, 0)
        uploaded_l.setSpacing(16)

        banner = QFrame()
        banner.setObjectName("allImportUploadedBanner")
        banner_l = QHBoxLayout(banner)
        banner_l.setContentsMargins(16, 16, 16, 16)
        banner_l.setSpacing(12)
        left_info = QHBoxLayout()
        left_info.setSpacing(12)
        file_icon = QLabel()
        file_icon.setFixedSize(24, 24)
        file_icon.setPixmap(_svg_pixmap(_FILE_SPREADSHEET_SVG, 24))
        left_info.addWidget(file_icon)
        name_col = QVBoxLayout()
        name_col.setSpacing(2)
        self._uploaded_name = QLabel("")
        self._uploaded_name.setObjectName("allImportUploadedName")
        name_col.addWidget(self._uploaded_name)
        self._uploaded_size = QLabel("")
        self._uploaded_size.setObjectName("allImportUploadedSize")
        name_col.addWidget(self._uploaded_size)
        left_info.addLayout(name_col)
        banner_l.addLayout(left_info, 1)
        pill = QFrame()
        pill.setObjectName("allImportSuccessPill")
        pill_l = QHBoxLayout(pill)
        pill_l.setContentsMargins(10, 4, 10, 4)
        pill_l.setSpacing(4)
        check_icon = QLabel()
        check_icon.setFixedSize(12, 12)
        check_icon.setPixmap(_svg_pixmap(_CHECK_GREEN_SVG, 12))
        pill_l.addWidget(check_icon)
        pill_text = QLabel("Uploaded")
        pill_text.setObjectName("allImportSuccessPillText")
        pill_l.addWidget(pill_text)
        banner_l.addWidget(pill)
        change_btn = QPushButton("Change file")
        change_btn.setObjectName("allImportChangeFile")
        change_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        change_btn.setToolTip("Remove this file and upload a different one")
        change_btn.clicked.connect(self._on_change_file)
        banner_l.addWidget(change_btn)
        uploaded_l.addWidget(banner)

        analysis_heading = QLabel("FILE ANALYSIS")
        analysis_heading.setObjectName("allImportAnalysisHeading")
        uploaded_l.addWidget(analysis_heading)
        analysis_badge = QFrame()
        analysis_badge.setObjectName("allImportAnalysisBadge")
        analysis_badge.setSizePolicy(QSizePolicy.Policy.Maximum, QSizePolicy.Policy.Fixed)
        badge_l = QHBoxLayout(analysis_badge)
        badge_l.setContentsMargins(10, 4, 10, 4)
        badge_text = QLabel("Analysis Complete")
        badge_text.setObjectName("allImportAnalysisBadgeText")
        badge_l.addWidget(badge_text)
        uploaded_l.addWidget(analysis_badge, 0, Qt.AlignmentFlag.AlignLeft)

        analysis_card = QFrame()
        analysis_card.setObjectName("allImportAnalysisCard")
        analysis_card_l = QVBoxLayout(analysis_card)
        analysis_card_l.setContentsMargins(12, 12, 12, 12)
        analysis_card_l.setSpacing(8)
        self._analysis_sheets_line = QLabel("")
        self._analysis_sheets_line.setObjectName("allImportAnalysisLine")
        self._analysis_sheets_line.setWordWrap(True)
        self._analysis_sheets_line.setTextFormat(Qt.TextFormat.RichText)
        analysis_card_l.addWidget(self._analysis_sheets_line)
        self._analysis_columns_line = QLabel("")
        self._analysis_columns_line.setObjectName("allImportAnalysisLine")
        self._analysis_columns_line.setWordWrap(True)
        analysis_card_l.addWidget(self._analysis_columns_line)
        uploaded_l.addWidget(analysis_card)
        uploaded_l.addStretch(1)
        self._uploaded_panel.setVisible(False)
        right_l.addWidget(self._uploaded_panel, 1)
        split.addWidget(right, 1)
        return page

    def _build_map_step(self) -> QWidget:
        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        summary = QFrame()
        summary.setObjectName("allImportFileSummary")
        summary.setFixedHeight(40)
        summary_l = QHBoxLayout(summary)
        summary_l.setContentsMargins(16, 10, 16, 10)
        summary_l.setSpacing(12)
        file_icon = QLabel()
        file_icon.setFixedSize(16, 16)
        file_icon.setPixmap(_svg_pixmap(_FILE_WHITE_SVG, 16))
        summary_l.addWidget(file_icon)
        self._map_summary_label = QLabel("")
        self._map_summary_label.setObjectName("allImportFileSummaryText")
        summary_l.addWidget(self._map_summary_label, 1)
        layout.addWidget(summary)

        controls = QHBoxLayout()
        controls.setSpacing(8)
        mode_group = QFrame()
        mode_group.setObjectName("allImportModeGroup")
        mode_l = QHBoxLayout(mode_group)
        mode_l.setContentsMargins(4, 4, 4, 4)
        mode_l.setSpacing(4)
        self.new_table_btn = QPushButton("New Table")
        self.new_table_btn.setObjectName("allImportModeBtn")
        self.new_table_btn.setCheckable(True)
        self.new_table_btn.setChecked(True)
        self.new_table_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.existing_table_btn = QPushButton("Existing Table")
        self.existing_table_btn.setObjectName("allImportModeBtn")
        self.existing_table_btn.setCheckable(True)
        self.existing_table_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        self.new_table_btn.clicked.connect(lambda: self._set_table_mode(True))
        self.existing_table_btn.clicked.connect(lambda: self._set_table_mode(False))
        mode_l.addWidget(self.new_table_btn)
        mode_l.addWidget(self.existing_table_btn)
        controls.addWidget(mode_group)
        controls.addStretch(1)
        self._sheets_bar = QHBoxLayout()
        self._sheets_bar.setSpacing(8)
        controls.addLayout(self._sheets_bar)
        layout.addLayout(controls)

        self.map_table = QTableWidget(0, 7)
        self.map_table.setObjectName("allImportMapTable")
        self.map_table.setHorizontalHeaderLabels(
            ["SELECT", "SOURCE COLUMN", "TARGET COLUMN", "DATA TYPE", "LENGTH", "NULLABLE", "PRIMARY KEY"]
        )
        apply_data_table_appearance(self.map_table, read_only=False, hide_vertical_header=True)
        self.map_table.setAlternatingRowColors(True)
        self.map_table.setShowGrid(False)
        self.map_table.setSortingEnabled(False)
        hh = self.map_table.horizontalHeader()
        hh.setObjectName("allImportMapHeader")
        hh.setSortIndicatorShown(False)
        hh.setFixedHeight(40)
        hh.setDefaultAlignment(Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter)
        vh = self.map_table.verticalHeader()
        vh.setDefaultSectionSize(40)
        vh.setVisible(False)
        # Re-apply map-specific chrome after shared data-table styles.
        self.map_table.setStyleSheet(
            f"""
            QTableWidget#allImportMapTable {{
                background: #ffffff; alternate-background-color: {_CARD_BG};
                border: 1px solid {_BORDER}; border-radius: 8px; gridline-color: {_BORDER};
                font-size: 13px; color: {_TEXT_DARK}; outline: none;
            }}
            QTableWidget#allImportMapTable::item {{
                padding: 4px 8px; border: none;
            }}
            QHeaderView#allImportMapHeader::section {{
                background: {_CARD_BG}; color: {_TEXT_MUTED}; font-size: 11px; font-weight: 700;
                border: none; border-bottom: 1px solid {_BORDER}; border-right: 1px solid {_BORDER};
                padding: 0 8px;
            }}
            """
        )
        hh.setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.map_table.setColumnWidth(0, 72)
        hh.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        hh.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        hh.setSectionResizeMode(3, QHeaderView.ResizeMode.Fixed)
        self.map_table.setColumnWidth(3, 140)
        hh.setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
        self.map_table.setColumnWidth(4, 90)
        hh.setSectionResizeMode(5, QHeaderView.ResizeMode.Fixed)
        self.map_table.setColumnWidth(5, 90)
        hh.setSectionResizeMode(6, QHeaderView.ResizeMode.Fixed)
        self.map_table.setColumnWidth(6, 110)
        attach_table_copy_shortcut(self.map_table)
        layout.addWidget(self.map_table, 1)
        return page

    def _build_execute_step(self) -> QWidget:
        self._execute_stack = QStackedWidget()
        self._execute_stack.addWidget(self._build_execute_settings_page())
        self._execute_stack.addWidget(self._build_success_page())
        return self._execute_stack

    def _build_execute_settings_page(self) -> QWidget:
        page = QWidget()
        split = QHBoxLayout(page)
        split.setContentsMargins(0, 0, 0, 0)
        split.setSpacing(24)

        left = QFrame()
        left.setObjectName("allImportLeftCard")
        left_l = QVBoxLayout(left)
        left_l.setContentsMargins(16, 16, 16, 16)
        left_l.setSpacing(16)
        left_title = QLabel("Execution Settings")
        left_title.setObjectName("allImportSectionTitle")
        left_l.addWidget(left_title)

        left_l.addWidget(self._field_caption("TARGET CONNECTION"))
        self.exec_connection_edit = QLineEdit()
        self.exec_connection_edit.setObjectName("allImportReadonlyField")
        self.exec_connection_edit.setReadOnly(True)
        self.exec_connection_edit.setFixedHeight(40)
        left_l.addWidget(self.exec_connection_edit)

        left_l.addWidget(self._field_caption("TARGET TABLE"))
        self.exec_table_edit = QLineEdit()
        self.exec_table_edit.setObjectName("allImportReadonlyField")
        self.exec_table_edit.setReadOnly(True)
        self.exec_table_edit.setFixedHeight(40)
        left_l.addWidget(self.exec_table_edit)

        left_l.addWidget(self._field_caption("TABLE OPERATION *"))
        op_row = QHBoxLayout()
        op_row.setSpacing(16)
        self.drop_create_radio = QRadioButton("DROP & CREATE")
        self.drop_create_radio.setObjectName("allImportOpRadio")
        self.drop_create_radio.setChecked(True)
        self.drop_create_radio.setCursor(Qt.CursorShape.PointingHandCursor)
        self.delete_radio = QRadioButton("DELETE")
        self.delete_radio.setObjectName("allImportOpRadio")
        self.delete_radio.setCursor(Qt.CursorShape.PointingHandCursor)
        self._op_group = QButtonGroup(self)
        self._op_group.addButton(self.drop_create_radio)
        self._op_group.addButton(self.delete_radio)
        self.drop_create_radio.toggled.connect(self._on_operation_toggled)
        self.delete_radio.toggled.connect(self._on_operation_toggled)
        op_row.addWidget(self.drop_create_radio)
        op_row.addWidget(self.delete_radio)
        op_row.addStretch(1)
        left_l.addLayout(op_row)
        left_l.addStretch(1)
        split.addWidget(left, 1)

        right = QFrame()
        right.setObjectName("allImportRightCard")
        right_l = QVBoxLayout(right)
        right_l.setContentsMargins(20, 20, 20, 20)
        right_l.setSpacing(20)
        right_title = QLabel("Import Progress")
        right_title.setObjectName("allImportSectionTitle")
        right_l.addWidget(right_title)

        progress_card = QFrame()
        progress_card.setObjectName("allImportProgressCard")
        progress_l = QVBoxLayout(progress_card)
        progress_l.setContentsMargins(16, 16, 16, 16)
        progress_l.setSpacing(12)

        status_row = QHBoxLayout()
        status_row.setSpacing(8)
        self._progress_spinner = QLabel()
        self._progress_spinner.setFixedSize(16, 16)
        self._progress_spinner.setPixmap(_svg_pixmap(_LOADER_SVG, 16))
        self._progress_spinner.setVisible(False)
        status_row.addWidget(self._progress_spinner)
        self._progress_status = QLabel("Ready to import")
        self._progress_status.setObjectName("allImportProgressStatus")
        status_row.addWidget(self._progress_status, 1)
        self._progress_pct = QLabel("0%")
        self._progress_pct.setObjectName("allImportProgressPct")
        status_row.addWidget(self._progress_pct)
        progress_l.addLayout(status_row)

        self._progress_bar = QProgressBar()
        self._progress_bar.setObjectName("allImportProgressBar")
        self._progress_bar.setRange(0, 100)
        self._progress_bar.setValue(0)
        self._progress_bar.setTextVisible(False)
        self._progress_bar.setFixedHeight(8)
        progress_l.addWidget(self._progress_bar)
        right_l.addWidget(progress_card)

        metrics = QVBoxLayout()
        metrics.setSpacing(10)
        self._metric_processed = self._metric_row(metrics, "ROWS PROCESSED", "0 / 0")
        self._metric_success = self._metric_row(
            metrics, "SUCCESSFUL", "0", value_object="allImportMetricSuccess"
        )
        self._metric_failed = self._metric_row(
            metrics, "FAILED", "0", value_object="allImportMetricFailed"
        )
        right_l.addLayout(metrics)
        right_l.addStretch(1)
        split.addWidget(right, 1)
        return page

    def _build_success_page(self) -> QWidget:
        page = QWidget()
        outer = QVBoxLayout(page)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setAlignment(Qt.AlignmentFlag.AlignCenter)

        panel = QWidget()
        layout = QVBoxLayout(panel)
        layout.setContentsMargins(40, 40, 40, 40)
        layout.setSpacing(24)
        layout.setAlignment(Qt.AlignmentFlag.AlignHCenter)

        hero = QVBoxLayout()
        hero.setSpacing(12)
        hero.setAlignment(Qt.AlignmentFlag.AlignHCenter)

        icon_wrap = QFrame()
        icon_wrap.setFixedSize(64, 64)
        icon_wrap.setStyleSheet(
            "QFrame { background: #dcfce7; border: none; border-radius: 32px; }"
        )
        icon_l = QVBoxLayout(icon_wrap)
        icon_l.setContentsMargins(0, 0, 0, 0)
        icon_l.setAlignment(Qt.AlignmentFlag.AlignCenter)
        icon = QLabel()
        icon.setFixedSize(32, 32)
        icon.setPixmap(_svg_pixmap(_CHECK_SUCCESS_SVG, 32))
        icon_l.addWidget(icon)
        hero.addWidget(icon_wrap, 0, Qt.AlignmentFlag.AlignHCenter)

        title = QLabel("Import Successful")
        title.setObjectName("allImportSuccessTitle")
        title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        hero.addWidget(title)

        subtitle = QLabel(
            "Your import session has executed. All detected data has been processed."
        )
        subtitle.setObjectName("allImportSuccessSubtitle")
        subtitle.setAlignment(Qt.AlignmentFlag.AlignCenter)
        subtitle.setWordWrap(True)
        hero.addWidget(subtitle)
        layout.addLayout(hero)

        card = QFrame()
        card.setObjectName("allImportSuccessSummary")
        card.setFixedWidth(500)
        card_l = QVBoxLayout(card)
        card_l.setContentsMargins(20, 20, 20, 20)
        card_l.setSpacing(16)

        heading = QLabel("SUMMARY DETAILS")
        heading.setObjectName("allImportSuccessSummaryHeading")
        card_l.addWidget(heading)

        self._success_operation = self._success_detail_row(card_l, "Operation Type", "DROP & CREATE")
        self._success_detail_row(card_l, "Engine Strategy", "Bulk Insertion")

        status_row = QHBoxLayout()
        status_row.setContentsMargins(0, 0, 0, 0)
        status_lbl = QLabel("Status")
        status_lbl.setObjectName("allImportSuccessRowLabel")
        status_row.addWidget(status_lbl)
        status_row.addStretch(1)
        status_badge = QLabel("200 OK")
        status_badge.setObjectName("allImportStatusBadge")
        status_row.addWidget(status_badge)
        card_l.addLayout(status_row)

        divider = QFrame()
        divider.setObjectName("allImportSuccessDivider")
        divider.setFrameShape(QFrame.Shape.HLine)
        card_l.addWidget(divider)

        self._success_total = self._success_detail_row(card_l, "Total Rows", "0")
        self._success_ok_rows = self._success_detail_row(
            card_l, "Successful Rows", "0", value_object="allImportSuccessRowSuccess"
        )
        self._success_fail_rows = self._success_detail_row(
            card_l, "Failed Rows", "0", value_object="allImportSuccessRowFailed"
        )
        layout.addWidget(card, 0, Qt.AlignmentFlag.AlignHCenter)

        actions = QHBoxLayout()
        actions.setSpacing(12)
        actions.setAlignment(Qt.AlignmentFlag.AlignHCenter)
        view_btn = QPushButton("View Details")
        view_btn.setObjectName("allImportViewDetails")
        view_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        view_btn.setFixedHeight(36)
        view_btn.clicked.connect(self._on_view_details)
        actions.addWidget(view_btn)
        back_btn = QPushButton("Back to Imports")
        back_btn.setObjectName("allImportBackToImports")
        back_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        back_btn.setFixedHeight(36)
        back_btn.clicked.connect(self._on_back_to_imports)
        actions.addWidget(back_btn)
        new_btn = QPushButton("New Import")
        new_btn.setObjectName("allImportNewImport")
        new_btn.setCursor(Qt.CursorShape.PointingHandCursor)
        new_btn.setFixedHeight(36)
        new_btn.clicked.connect(self._on_new_import)
        actions.addWidget(new_btn)
        layout.addLayout(actions)

        outer.addWidget(panel)
        return page

    def _success_detail_row(
        self,
        parent_layout: QVBoxLayout,
        label: str,
        value: str,
        *,
        value_object: str = "allImportSuccessRowValue",
    ) -> QLabel:
        row = QHBoxLayout()
        row.setContentsMargins(0, 0, 0, 0)
        lbl = QLabel(label)
        lbl.setObjectName("allImportSuccessRowLabel")
        row.addWidget(lbl)
        row.addStretch(1)
        val = QLabel(value)
        val.setObjectName(value_object)
        row.addWidget(val)
        parent_layout.addLayout(row)
        return val

    def _metric_row(
        self,
        parent_layout: QVBoxLayout,
        label: str,
        value: str,
        *,
        value_object: str = "allImportMetricValue",
    ) -> QLabel:
        row = QHBoxLayout()
        row.setContentsMargins(0, 0, 0, 0)
        lbl = QLabel(label)
        lbl.setObjectName("allImportMetricLabel")
        row.addWidget(lbl)
        row.addStretch(1)
        val = QLabel(value)
        val.setObjectName(value_object)
        row.addWidget(val)
        parent_layout.addLayout(row)
        return val

    @staticmethod
    def _field_caption(text: str) -> QLabel:
        lbl = QLabel(text)
        lbl.setObjectName("allImportFieldCaption")
        return lbl

    def go_to_nav_item(self, _nav_item: str) -> None:
        return

    def showEvent(self, event) -> None:  # noqa: N802
        super().showEvent(event)
        if not self._ctx.connections():
            token = self._token()
            if token:
                self._ctx.refresh_connections(token)

    def _refresh_connections(self) -> None:
        current = self.connection_combo.currentData()
        current_id = None
        if isinstance(current, dict):
            current_id = current.get("connectionId") or current.get("id")
        self.connection_combo.blockSignals(True)
        self.connection_combo.clear()
        self.connection_combo.addItem("Select connection", None)
        for conn in self._ctx.connections():
            if isinstance(conn, dict):
                self.connection_combo.addItem(etl_connection_display_label(conn), conn)
        if current_id is not None:
            for i in range(self.connection_combo.count()):
                data = self.connection_combo.itemData(i)
                if isinstance(data, dict) and str(data.get("connectionId") or data.get("id")) == str(current_id):
                    self.connection_combo.setCurrentIndex(i)
                    break
        else:
            self.connection_combo.setCurrentIndex(0)
        self.connection_combo.blockSignals(False)

    def _on_tab_clicked(self, index: int) -> None:
        has_data = bool(self._selected_path or (self._remote_session and self._sheets))
        already_analyzed = bool(self._session_id and self._sheets)
        if index == 1 and not has_data:
            self._show_error("Select and analyze a file before mapping columns.")
            self._apply_step_chrome()
            return
        if index == 1 and not already_analyzed:
            # Do not create/upload from the tab — only Next runs that pipeline.
            self._show_error("Click Next to upload and analyze before mapping columns.")
            self._apply_step_chrome()
            return
        if index == 2 and self._step_index < 1 and not (
            self._remote_session and self._sheets
        ):
            self._show_error("Complete Map Columns before Execute & Result.")
            self._apply_step_chrome()
            return
        if index == 1:
            self._go_to_map_columns(reuse_session=True)
            return
        if index == 2:
            if self._step_index >= 2:
                self._set_step(2)
                return
            self._go_to_execute()
            return
        self._set_step(index)

    def _set_step(self, index: int) -> None:
        self._step_index = index
        self.stack.setCurrentIndex(index)
        self._apply_step_chrome()

    def _apply_step_chrome(self) -> None:
        for i, btn in enumerate(self._tab_buttons):
            active = i == self._step_index
            completed = i < self._step_index
            btn.setChecked(active)
            label = self._tab_step_labels[i]
            if completed:
                btn.setText(f"  ✓   {label}")
            else:
                btn.setText(f"  {i + 1}   {label}")
        showing_success = (
            self._step_index == 2
            and getattr(self, "_execute_stack", None) is not None
            and self._execute_stack.currentIndex() == 1
        )
        if showing_success:
            self.primary_btn.setVisible(False)
            self.secondary_btn.setVisible(False)
            return
        self.primary_btn.setVisible(True)
        self.secondary_btn.setVisible(True)
        if self._step_index == 0:
            can_next = bool(self._selected_path or (self._remote_session and self._session_id))
            self.primary_btn.setEnabled(True)
            self.primary_btn.setText("Next" if can_next else "Import")
            self.secondary_btn.setText("Cancel")
        elif self._step_index == 1:
            self.primary_btn.setEnabled(True)
            self.primary_btn.setText("Next")
            self.secondary_btn.setText("Back")
        elif self._importing:
            self.primary_btn.setEnabled(False)
            self.primary_btn.setText("Importing...")
            self.secondary_btn.setText("Cancel")
        else:
            self.primary_btn.setEnabled(True)
            self.primary_btn.setText("Execute")
            self.secondary_btn.setText("Back")

    def _token(self) -> str | None:
        profile = get_user_profile() or {}
        token = (
            profile.get("token")
            or profile.get("accessToken")
            or profile.get("access_token")
            or profile.get("jwt")
        )
        return str(token) if token else None

    def _set_busy(self, busy: bool, *, status: str = "") -> None:
        self._busy = busy
        self.primary_btn.setEnabled(not busy)
        self.secondary_btn.setEnabled(not busy)
        if busy and status:
            show_page_info(self, self.message_label, status, sticky=True)
        QApplication.processEvents()

    def _go_to_execute(self) -> None:
        body, err = self._build_mapping_body()
        if err or body is None:
            self._show_error(err or "Could not build mapping.")
            return
        if not self._session_id:
            self._show_error("Import session is missing. Re-run Upload & Analyze.")
            return
        self._set_busy(True, status="Saving column mapping…")
        try:
            save = api_save_import_mapping(self._session_id, body, token=self._token())
        finally:
            self._set_busy(False)
        if not save.get("success"):
            self._show_error(str(save.get("message") or "Failed to save mapping."))
            return
        # Validate selected operation while still on Map Columns so errors stay here.
        self.drop_create_radio.blockSignals(True)
        self.delete_radio.blockSignals(True)
        self.drop_create_radio.setChecked(True)
        self.drop_create_radio.blockSignals(False)
        self.delete_radio.blockSignals(False)
        if not self._run_operation_validate(show_success=False):
            return
        conn = self.connection_combo.currentData()
        if isinstance(conn, dict):
            self.exec_connection_edit.setText(etl_connection_display_label(conn))
        else:
            self.exec_connection_edit.setText(self.connection_combo.currentText())
        self.exec_table_edit.setText(self.target_table_edit.text().strip())
        self._reset_import_progress()
        self._execute_stack.setCurrentIndex(0)
        self._clear_message()
        self._set_step(2)

    def _selected_operation(self) -> str:
        return "DROP_CREATE" if self.drop_create_radio.isChecked() else "DELETE"

    def _active_sheet_id(self) -> int | str | None:
        sheet = self._sheets[self._active_sheet] if self._sheets else {}
        return _first_id(sheet, "sheetId", "sheet_id", "sheetID", "id")

    def _on_operation_toggled(self, checked: bool) -> None:
        if not checked:
            return
        if self._step_index != 2 or self._importing or self._import_complete:
            return
        if self._execute_stack.currentIndex() != 0:
            return
        self._run_operation_validate(show_success=True)

    def _run_operation_validate(self, *, show_success: bool = False) -> bool:
        """POST ``imports/{sessionId}/validate`` for the selected DROP_CREATE/DELETE."""
        if not self._session_id:
            self._show_error("Import session is missing. Re-run Upload & Analyze.")
            return False
        sheet_id = self._active_sheet_id()
        if sheet_id is None:
            self._show_error("Sheet ID is missing. Re-run analyze.")
            return False
        operation = self._selected_operation()
        self._set_busy(True, status=f"Validating {operation.replace('_', ' ').title()}…")
        try:
            result = api_validate_import_sheet(
                self._session_id,
                sheet_id,
                operation=operation,
                token=self._token(),
            )
        finally:
            self._set_busy(False)
        if not result.get("success"):
            self._show_error(str(result.get("message") or "Validation failed."))
            return False
        if show_success:
            show_page_success(
                self,
                self.message_label,
                str(result.get("message") or "Validation passed."),
            )
        else:
            self._clear_message()
        return True

    def _reset_import_progress(self) -> None:
        self._progress_timer.stop()
        self._importing = False
        self._import_complete = False
        self._progress_value = 0
        self._total_rows = 0
        self._progress_spinner.setVisible(False)
        self._progress_status.setText("Ready to import")
        self._progress_pct.setText("0%")
        self._progress_bar.setValue(0)
        self._metric_processed.setText("0 / 0")
        self._metric_success.setText("0")
        self._metric_failed.setText("0")

    def _start_import(self) -> None:
        if self._importing:
            return
        if not self._session_id:
            self._show_error("Import session is missing. Re-run Upload & Analyze.")
            return
        sheet_id = self._active_sheet_id()
        if sheet_id is None:
            self._show_error("Sheet ID is missing. Re-run analyze.")
            return
        if not self._run_operation_validate(show_success=False):
            return
        operation = self._selected_operation()
        self._importing = True
        self._import_complete = False
        self._progress_value = 0
        self._total_rows = 0
        self._execute_stack.setCurrentIndex(0)
        self._progress_spinner.setVisible(True)
        self._progress_status.setText("Importing...")
        self._progress_pct.setText("0%")
        self._progress_bar.setValue(10)
        self._metric_processed.setText("—")
        self._metric_success.setText("—")
        self._metric_failed.setText("—")
        self._clear_message()
        self._apply_step_chrome()
        QApplication.processEvents()

        result = api_execute_import_sheet(
            self._session_id,
            sheet_id,
            operation=operation,
            token=self._token(),
        )
        self._progress_bar.setValue(100)
        self._progress_pct.setText("100%")
        self._progress_spinner.setVisible(False)
        self._importing = False
        if not result.get("success"):
            self._progress_status.setText("Import failed")
            self._apply_step_chrome()
            self._show_error(str(result.get("message") or "Execute/import failed."))
            return

        data = result.get("data") if isinstance(result.get("data"), dict) else {}
        raw = result if isinstance(result, dict) else {}
        total = (
            data.get("totalRows")
            or data.get("rowCount")
            or data.get("total")
            or raw.get("totalRows")
            or 0
        )
        success = (
            data.get("successfulRows")
            or data.get("successCount")
            or data.get("successRows")
            or total
            or 0
        )
        failed = (
            data.get("failedRows")
            or data.get("failCount")
            or data.get("failedCount")
            or 0
        )
        try:
            total_i = int(total or 0)
            success_i = int(success or 0)
            failed_i = int(failed or 0)
        except (TypeError, ValueError):
            total_i, success_i, failed_i = 0, 0, 0
        if total_i <= 0 and success_i > 0:
            total_i = success_i + failed_i
        self._total_rows = total_i
        self._import_complete = True
        self._progress_status.setText("Import complete")
        self._metric_processed.setText(f"{total_i:,} / {total_i:,}" if total_i else "—")
        self._metric_success.setText(f"{success_i:,}" if success_i or total_i else "—")
        self._metric_failed.setText(f"{failed_i:,}")

        op_label = "DROP & CREATE" if self._selected_operation() == "DROP_CREATE" else "DELETE"
        self._success_operation.setText(op_label)
        self._success_total.setText(f"{total_i:,}" if total_i else "—")
        self._success_ok_rows.setText(f"{success_i:,}" if success_i or total_i else "—")
        self._success_fail_rows.setText(f"{failed_i:,}")
        self._clear_message()
        self._execute_stack.setCurrentIndex(1)
        self._apply_step_chrome()

    def _tick_import_progress(self) -> None:
        # Kept for compatibility; real execute is synchronous via API.
        if self._progress_value >= 100:
            self._progress_timer.stop()
            return
        self._progress_value = min(100, self._progress_value + 5)
        self._progress_bar.setValue(self._progress_value)
        self._progress_pct.setText(f"{self._progress_value}%")

    def _finish_import(self) -> None:
        self._progress_timer.stop()

    def _cancel_import(self) -> None:
        self._progress_timer.stop()
        self._importing = False
        self._reset_import_progress()
        self._execute_stack.setCurrentIndex(0)
        self._apply_step_chrome()
        self._clear_message()

    def _on_view_details(self) -> None:
        self._execute_stack.setCurrentIndex(0)
        self._apply_step_chrome()
        if self._import_complete:
            self.primary_btn.setText("Done")
            self.primary_btn.setEnabled(True)
            self.secondary_btn.setText("Back")

    def reset_for_new_import(self) -> None:
        """Clear wizard state without emitting exit-to-list (used by hub New Import)."""
        self.import_name_edit.clear()
        self.file_type_combo.setCurrentIndex(0)
        self.encoding_combo.setCurrentIndex(0)
        self.connection_combo.setCurrentIndex(0)
        self.target_table_edit.clear()
        self._clear_uploaded_state()
        self._session_id = ""
        self._sheets = []
        self._remote_session = False
        self._import_complete = False
        self._reset_import_progress()
        self._execute_stack.setCurrentIndex(0)
        self._clear_message()
        self._set_step(0)

    def load_existing_session(self, row: dict[str, Any]) -> None:
        """Open wizard pre-filled from an Import Sessions list row."""
        self.reset_for_new_import()
        sid = str(row.get("sessionId") or "").strip()
        if not sid:
            self._show_error("This import session has no session ID.")
            return

        token = self._token()
        if not token:
            self._show_error("Session expired. Please log in again.")
            return

        self._remote_session = True
        self._session_id = sid
        name = str(row.get("sessionName") or sid).strip()
        self.import_name_edit.setText(name)

        source_type = str(row.get("sourceType") or row.get("fileType") or "").strip()
        self._select_file_type(source_type)
        self._select_encoding(str(row.get("encoding") or ""))

        table = str(
            row.get("targetTableName") or row.get("tableName") or ""
        ).strip()
        if table:
            self.target_table_edit.setText(table)

        conn_hint = (
            row.get("targetConnectionId")
            or row.get("connectionId")
            or row.get("targetConnectionName")
            or row.get("connectionName")
        )
        self._select_connection(conn_hint)

        self._set_busy(True, status="Loading session…")
        try:
            present = api_analyze_import_if_present(sid, token=token)
            listed = api_get_import_sheets_by_uuid(sid, token=token)
            sheets: list[dict[str, Any]] = []
            if present.get("success") and present.get("present"):
                sheets = [s for s in (present.get("data") or []) if isinstance(s, dict)]
                raw = present.get("raw") if isinstance(present.get("raw"), dict) else {}
                file_name = str(
                    raw.get("fileName")
                    or row.get("fileName")
                    or name
                ).strip() or name
                if present.get("targetTableName") and not self.target_table_edit.text().strip():
                    self.target_table_edit.setText(str(present.get("targetTableName")))
                if present.get("targetConnectionId") is not None:
                    self._select_connection(present.get("targetConnectionId"))
                if raw.get("sourceType") or raw.get("fileType"):
                    self._select_file_type(
                        str(raw.get("sourceType") or raw.get("fileType") or "")
                    )
                if raw.get("encoding"):
                    self._select_encoding(str(raw.get("encoding")))
                self._show_remote_uploaded_state(file_name, sheets)
            elif listed.get("success"):
                sheets = [s for s in (listed.get("data") or []) if isinstance(s, dict)]
                file_name = str(row.get("fileName") or name).strip() or name
                if sheets:
                    self._show_remote_uploaded_state(file_name, sheets)

            if listed.get("success"):
                listed_sheets = [
                    s for s in (listed.get("data") or []) if isinstance(s, dict)
                ]
                if sheets and listed_sheets:
                    sheets = self._merge_sheet_ids(sheets, listed_sheets)
                elif listed_sheets and not sheets:
                    sheets = listed_sheets

            self._sheets = sheets
            self._active_sheet = 0
            if sheets:
                names = [str(s.get("name") or f"Sheet{i + 1}") for i, s in enumerate(sheets)]
                first_cols = (
                    sheets[0].get("columns")
                    if isinstance(sheets[0].get("columns"), list)
                    else []
                )
                self._analysis_sheets_line.setText(
                    f"• {len(names)} sheet{'s' if len(names) != 1 else ''} discovered: "
                    f"<span style='font-weight:600;'>{', '.join(names)}</span>"
                )
                self._analysis_columns_line.setText(
                    f"• {len(first_cols)} column{'s' if len(first_cols) != 1 else ''} "
                    f"detected in {names[0]}"
                )
                # Prefer table name from first sheet when form still empty.
                if not self.target_table_edit.text().strip():
                    for key in ("tableName", "targetTableName"):
                        value = sheets[0].get(key)
                        if value is not None and str(value).strip():
                            self.target_table_edit.setText(str(value).strip())
                            break

            # Always land on Upload & Analyze; user advances via Next.
            self._set_step(0)
            self._clear_message()
        finally:
            self._set_busy(False)

    def _select_file_type(self, raw: str) -> None:
        text = str(raw or "").strip().upper().replace("-", "_")
        if not text:
            return
        mapping = {
            "EXCEL": "Excel",
            "XLSX": "Excel",
            "XLS": "Excel",
            "CSV": "CSV",
            "JSON": "JSON",
        }
        label = mapping.get(text, text.title())
        idx = self.file_type_combo.findText(label)
        if idx >= 0:
            self.file_type_combo.setCurrentIndex(idx)

    def _select_encoding(self, raw: str) -> None:
        text = str(raw or "").strip().upper().replace("-", "_").replace(" ", "_")
        if not text:
            return
        for i in range(self.encoding_combo.count()):
            data = str(self.encoding_combo.itemData(i) or "").strip().upper()
            label = self.encoding_combo.itemText(i).upper().replace("-", "_").replace(" ", "_")
            if text == data or text in label or data in text:
                self.encoding_combo.setCurrentIndex(i)
                return

    def _select_connection(self, hint: Any) -> None:
        if hint is None or str(hint).strip() == "":
            return
        if not self._ctx.connections():
            token = self._token()
            if token:
                self._ctx.refresh_connections(token)
                self._refresh_connections()
        want_id = None
        want_name = None
        if isinstance(hint, dict):
            want_id = hint.get("connectionId") or hint.get("id")
            want_name = (
                hint.get("connectionName")
                or hint.get("name")
                or hint.get("label")
            )
        else:
            text = str(hint).strip()
            if text.isdigit():
                want_id = text
            else:
                want_name = text
        for i in range(self.connection_combo.count()):
            data = self.connection_combo.itemData(i)
            if not isinstance(data, dict):
                continue
            cid = str(data.get("connectionId") or data.get("id") or "").strip()
            label = etl_connection_display_label(data)
            name = str(
                data.get("connectionName") or data.get("name") or label or ""
            ).strip()
            if want_id is not None and cid == str(want_id).strip():
                self.connection_combo.setCurrentIndex(i)
                return
            if want_name and want_name.lower() in {
                name.lower(),
                label.lower(),
            }:
                self.connection_combo.setCurrentIndex(i)
                return

    def _show_remote_uploaded_state(
        self, file_name: str, sheets: list[dict[str, Any]]
    ) -> None:
        self._selected_path = ""
        self._file_size_text = "Server file"
        self._uploaded_name.setText(file_name or "Uploaded file")
        self._uploaded_size.setText(self._file_size_text)
        names = [str(s.get("name") or f"Sheet{i + 1}") for i, s in enumerate(sheets)]
        first_cols = (
            sheets[0].get("columns")
            if sheets and isinstance(sheets[0].get("columns"), list)
            else []
        )
        if names:
            self._analysis_sheets_line.setText(
                f"• {len(names)} sheet{'s' if len(names) != 1 else ''} discovered: "
                f"<span style='font-weight:600;'>{', '.join(names)}</span>"
            )
            self._analysis_columns_line.setText(
                f"• {len(first_cols)} column{'s' if len(first_cols) != 1 else ''} "
                f"detected in {names[0]}"
            )
        self.drop_zone.setVisible(False)
        self._uploaded_panel.setVisible(True)
        self._apply_step_chrome()

    def _open_map_from_loaded_session(self) -> None:
        file_name = self._uploaded_name.text().strip() or "session"
        file_type = self.file_type_combo.currentText() or "Excel"
        sheet_count = len(self._sheets)
        self._map_summary_label.setText(
            f"{file_name} · {file_type} · {sheet_count} sheet{'s' if sheet_count != 1 else ''} · {self._file_size_text}"
        )
        self._active_sheet = 0
        self._rebuild_sheet_buttons()
        self._populate_map_table()
        self._set_step(1)

    def _on_new_import(self) -> None:
        """Start another import wizard from the success screen."""
        self.reset_for_new_import()

    def _on_back_to_imports(self) -> None:
        """Return to the Excel All Import dashboard table."""
        self.reset_for_new_import()
        self.exit_to_list_requested.emit()

    def _on_file_selected(self, path: str) -> None:
        self._selected_path = path
        self._remote_session = False
        self._session_id = ""
        self._sheets = []
        self._clear_message()
        self._show_uploaded_state(path)

    def _on_change_file(self) -> None:
        """Clear the selected file so the user can upload a different one."""
        if self._busy or self._importing:
            return
        self._clear_uploaded_state()
        self._clear_message()
        self._apply_step_chrome()

    def _show_uploaded_state(self, path: str) -> None:
        file_path = Path(path)
        self._file_size_text = _format_file_size(file_path.stat().st_size) if file_path.is_file() else "—"
        self._uploaded_name.setText(file_path.name)
        self._uploaded_size.setText(self._file_size_text)
        self._preview_sheets = _analyze_workbook(path)
        if self._preview_sheets:
            names = [str(s.get("name") or "") for s in self._preview_sheets]
            first_headers = list(self._preview_sheets[0].get("headers") or [])
            self._analysis_sheets_line.setText(
                f"• {len(names)} sheet{'s' if len(names) != 1 else ''} discovered: "
                f"<span style='font-weight:600;'>{', '.join(names)}</span>"
            )
            self._analysis_columns_line.setText(
                f"• {len(first_headers)} column{'s' if len(first_headers) != 1 else ''} "
                f"detected in {names[0]}"
            )
        else:
            self._analysis_sheets_line.setText("• File selected — click Next to upload & analyze")
            self._analysis_columns_line.setText("• Column detection runs on the server")
        self.drop_zone.setVisible(False)
        self._uploaded_panel.setVisible(True)
        self._apply_step_chrome()

    def _clear_uploaded_state(self) -> None:
        self._selected_path = ""
        self._file_size_text = ""
        self._session_id = ""
        self._sheets = []
        self._preview_sheets = []
        self._active_sheet = 0
        self._remote_session = False
        self._uploaded_name.clear()
        self._uploaded_size.clear()
        self._analysis_sheets_line.clear()
        self._analysis_columns_line.clear()
        self._uploaded_panel.setVisible(False)
        self.drop_zone.setVisible(True)
        self.map_table.setRowCount(0)
        self._rebuild_sheet_buttons()

    def _set_table_mode(self, new_table: bool) -> None:
        self.new_table_btn.setChecked(new_table)
        self.existing_table_btn.setChecked(not new_table)

    def _rebuild_sheet_buttons(self) -> None:
        while self._sheets_bar.count():
            item = self._sheets_bar.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()
        self._sheet_buttons = []
        for index, sheet in enumerate(self._sheets):
            name = str(sheet.get("name") or f"Sheet{index + 1}")
            btn = QPushButton(name)
            btn.setObjectName("allImportSheetBtn")
            btn.setCheckable(True)
            btn.setChecked(index == self._active_sheet)
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.clicked.connect(lambda _c=False, i=index: self._select_sheet(i))
            self._sheets_bar.addWidget(btn)
            self._sheet_buttons.append(btn)

    def _select_sheet(self, index: int) -> None:
        self._active_sheet = index
        for i, btn in enumerate(self._sheet_buttons):
            btn.setChecked(i == index)
        self._populate_map_table()

    def _merge_sheet_ids(
        self,
        analyzed: list[dict[str, Any]],
        listed: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        by_name: dict[str, dict[str, Any]] = {}
        for row in listed:
            if not isinstance(row, dict):
                continue
            name = str(row.get("sheetName") or row.get("name") or "").strip().lower()
            if name:
                by_name[name] = row
        for sheet in analyzed:
            if not isinstance(sheet, dict):
                continue
            if sheet.get("sheetId") is not None:
                continue
            name = str(sheet.get("name") or "").strip().lower()
            match = by_name.get(name)
            if not match:
                continue
            sid = match.get("sheetId") or match.get("sheet_id") or match.get("id")
            if sid is not None:
                sheet["sheetId"] = sid
        return analyzed

    def _run_upload_analyze_pipeline(self) -> bool:
        """Create session → upload → analyze → get sheets. Returns True on success."""
        token = self._token()
        if not token:
            self._show_error("Session expired. Please log in again.")
            return False
        name = self.import_name_edit.text().strip()
        source_type = str(self.file_type_combo.currentData() or "EXCEL")
        encoding = str(self.encoding_combo.currentData() or "UTF_8")
        codepage = _encoding_codepage(encoding)

        self._set_busy(True, status="Creating import session…")
        try:
            create = api_create_import(
                name, source_type, encoding=encoding, token=token
            )
            if not create.get("success"):
                self._show_error(str(create.get("message") or "Failed to create import."))
                return False
            session_id = _session_id_from_create(create)
            if not session_id:
                self._show_error("Import created but session ID was not returned.")
                return False
            self._session_id = session_id

            self._set_busy(True, status="Uploading file…")
            upload = api_upload_import_file(
                session_id,
                self._selected_path,
                token=token,
                codepage=codepage,
                multi_language=True,
            )
            if not upload.get("success"):
                self._show_error(str(upload.get("message") or "Upload failed."))
                return False

            self._set_busy(True, status="Analyzing workbook…")
            analyze = api_analyze_import_file(
                session_id,
                token=token,
                codepage=codepage,
                multi_language=True,
            )
            if not analyze.get("success"):
                self._show_error(str(analyze.get("message") or "Analyze failed."))
                return False
            sheets = analyze.get("data") if isinstance(analyze.get("data"), list) else []
            sheets = [s for s in sheets if isinstance(s, dict)]

            self._set_busy(True, status="Loading sheets…")
            listed = api_get_import_sheets_by_uuid(session_id, token=token)
            listed_sheets = (
                listed.get("data") if isinstance(listed.get("data"), list) else []
            )
            if sheets:
                sheets = self._merge_sheet_ids(sheets, listed_sheets)
            elif listed.get("success") and listed_sheets:
                # Fallback: sheets endpoint only — no column metadata yet.
                sheets = []
                for row in listed_sheets:
                    if not isinstance(row, dict):
                        continue
                    sheets.append(
                        {
                            "name": str(row.get("sheetName") or row.get("name") or "Sheet"),
                            "sheetId": row.get("sheetId") or row.get("id"),
                            "columns": [],
                            "_raw": row,
                        }
                    )

            if not sheets:
                self._show_error("Analyze returned no sheets.")
                return False
            self._sheets = sheets

            names = [str(s.get("name") or "") for s in sheets]
            first_cols = sheets[0].get("columns") if isinstance(sheets[0].get("columns"), list) else []
            self._analysis_sheets_line.setText(
                f"• {len(names)} sheet{'s' if len(names) != 1 else ''} discovered: "
                f"<span style='font-weight:600;'>{', '.join(names)}</span>"
            )
            self._analysis_columns_line.setText(
                f"• {len(first_cols)} column{'s' if len(first_cols) != 1 else ''} "
                f"detected in {names[0]}"
            )
            return True
        finally:
            self._set_busy(False)

    def _go_to_map_columns(self, *, reuse_session: bool = False) -> None:
        has_remote = bool(self._remote_session and self._session_id and self._sheets)
        if not self._selected_path and not has_remote:
            self._show_error("Select a file first.")
            return
        if self._selected_path and not self._validate_upload_form():
            return
        if has_remote and not self._selected_path:
            self._open_map_from_loaded_session()
            self._clear_message()
            return
        already = bool(self._session_id and self._sheets)
        if reuse_session and not already:
            self._show_error("Click Next to upload and analyze before mapping columns.")
            return
        if not already:
            if not self._run_upload_analyze_pipeline():
                return
        file_name = Path(self._selected_path).name
        file_type = self.file_type_combo.currentText() or "Excel"
        sheet_count = len(self._sheets)
        self._map_summary_label.setText(
            f"{file_name} · {file_type} · {sheet_count} sheet{'s' if sheet_count != 1 else ''} · {self._file_size_text}"
        )
        if not (reuse_session and already):
            self._active_sheet = 0
        self._rebuild_sheet_buttons()
        self._populate_map_table()
        self._clear_message()
        self._set_step(1)

    def _populate_map_table(self) -> None:
        self.map_table.setRowCount(0)
        if not self._sheets:
            return
        sheet = self._sheets[self._active_sheet]
        columns = sheet.get("columns") if isinstance(sheet.get("columns"), list) else []
        if not columns:
            # Fallback to local preview headers if API returned no columns.
            preview = (
                self._preview_sheets[self._active_sheet]
                if self._active_sheet < len(self._preview_sheets)
                else {}
            )
            headers = list(preview.get("headers") or [])
            columns = [
                {
                    "name": h,
                    "tgtColumnName": _to_snake_case(str(h)),
                    "sourceColumnId": None,
                    "selected": True,
                    "nullableBool": True,
                    "primaryKey": i == 0,
                }
                for i, h in enumerate(headers)
            ]
        self.map_table.setRowCount(len(columns))
        for row, col in enumerate(columns):
            if not isinstance(col, dict):
                continue
            source_name = str(col.get("name") or f"column_{row + 1}")
            selected = bool(col.get("selected", True))
            self.map_table.setRowHeight(row, 40)

            select_box = QCheckBox()
            select_box.setChecked(selected)
            select_box.setProperty("sourceColumnId", col.get("sourceColumnId"))
            select_box.setStyleSheet("background: transparent; border: none;")
            self.map_table.setCellWidget(row, 0, _map_centered_widget(select_box))

            source_item = QTableWidgetItem(source_name)
            source_item.setFlags(source_item.flags() & ~Qt.ItemFlag.ItemIsEditable)
            self.map_table.setItem(row, 1, source_item)

            target_name = _bracket_spaced_identifier(
                str(col.get("tgtColumnName") or _to_snake_case(source_name))
            )
            target = QLineEdit(target_name)
            target.setFixedHeight(28)
            target.setStyleSheet(
                f"QLineEdit {{ background: transparent; border: 1px solid {_BORDER_INPUT}; "
                "border-radius: 4px; padding: 2px 8px; font-size: 13px; }"
            )
            target.editingFinished.connect(
                lambda edit=target: edit.setText(_bracket_spaced_identifier(edit.text()))
            )
            self.map_table.setCellWidget(row, 2, target)

            api_dtype = str(col.get("tgtDatatype") or col.get("datatype") or "").strip()
            dtype, length = _guess_datatype(source_name)
            if api_dtype:
                # Prefer SQL-ish UI label when API sends SQL or Java type.
                java = import_target_data_type_to_java(api_dtype)
                reverse = {
                    "INTEGER": "INT",
                    "LONG": "BIGINT",
                    "STRING": "NVARCHAR",
                    "BOOLEAN": "BIT",
                    "FLOAT": "FLOAT",
                    "DECIMAL": "DECIMAL",
                    "DATE": "DATE",
                    "DATETIME": "DATETIME",
                    "TIMESTAMP": "DATETIME",
                }
                dtype = reverse.get(java, api_dtype.upper().split("(", 1)[0] or dtype)
            length_raw = str(col.get("length") or "").strip()
            if length_raw:
                length = length_raw

            dtype_combo = QComboBox()
            for opt in _DATATYPE_OPTIONS:
                dtype_combo.addItem(opt)
            idx = dtype_combo.findText(dtype)
            if idx < 0:
                for i in range(dtype_combo.count()):
                    if dtype_combo.itemText(i).upper() == dtype.upper():
                        idx = i
                        break
            dtype_combo.setCurrentIndex(idx if idx >= 0 else 0)
            apply_form_combobox_field(dtype_combo, height_px=28, min_width=100)
            self.map_table.setCellWidget(row, 3, dtype_combo)

            length_item = QTableWidgetItem(length if length else "-")
            length_item.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            self.map_table.setItem(row, 4, length_item)

            nullable = QCheckBox()
            nullable.setChecked(bool(col.get("nullableBool", row > 0)))
            nullable.setStyleSheet("background: transparent; border: none;")
            self.map_table.setCellWidget(row, 5, _map_centered_widget(nullable))

            is_pk = bool(col.get("primaryKey", row == 0))
            pk = QLabel()
            pk.setAlignment(Qt.AlignmentFlag.AlignCenter)
            pk.setProperty("isPrimaryKey", is_pk)
            pk.setStyleSheet("background: transparent; border: none;")
            if is_pk:
                pk.setPixmap(_svg_pixmap(_KEY_BLUE_SVG, 14))
            else:
                pk.setText("—")
                pk.setStyleSheet(
                    f"color: {_TEXT_MUTED}; font-size: 13px; background: transparent; border: none;"
                )
            self.map_table.setCellWidget(row, 6, _map_centered_widget(pk))

    def _build_mapping_body(self) -> tuple[dict[str, Any] | None, str | None]:
        target_table = self.target_table_edit.text().strip()
        if not target_table:
            return None, "Enter a target table name before continuing."
        conn = self.connection_combo.currentData()
        conn_id = etl_connection_id(conn) if isinstance(conn, dict) else None
        if conn_id is None or str(conn_id).strip() == "":
            return None, "Select a connection before continuing."
        try:
            target_connection_id = int(conn_id)
        except (TypeError, ValueError):
            return None, "Connection ID must be numeric."
        if not self._sheets:
            return None, "Analyze a workbook before saving the mapping."
        sheet = self._sheets[self._active_sheet]
        sheet_id = _first_id(sheet, "sheetId", "sheet_id", "sheetID", "id")
        try:
            sheet_id_int = int(sheet_id) if sheet_id is not None and str(sheet_id).strip() else None
        except (TypeError, ValueError):
            sheet_id_int = None
        if sheet_id_int is None:
            return None, "Sheet ID is missing from analyze result."

        api_columns = sheet.get("columns") if isinstance(sheet.get("columns"), list) else []
        mappings: list[dict[str, Any]] = []
        for row in range(self.map_table.rowCount()):
            select_wrap = self.map_table.cellWidget(row, 0)
            select_box = (
                select_wrap.findChild(QCheckBox)
                if isinstance(select_wrap, QWidget)
                else None
            )
            selected = bool(select_box.isChecked()) if select_box is not None else True
            source_column_id = (
                select_box.property("sourceColumnId") if select_box is not None else None
            )
            if source_column_id is None and row < len(api_columns) and isinstance(api_columns[row], dict):
                source_column_id = api_columns[row].get("sourceColumnId")
            try:
                source_id_int = (
                    int(source_column_id)
                    if source_column_id is not None and str(source_column_id).strip()
                    else None
                )
            except (TypeError, ValueError):
                source_id_int = None
            if source_id_int is None:
                return None, "One or more columns are missing sourceColumnId from analyze."

            target_widget = self.map_table.cellWidget(row, 2)
            target_column_name = (
                target_widget.text().strip()
                if isinstance(target_widget, QLineEdit)
                else ""
            )
            target_column_name = _bracket_spaced_identifier(target_column_name)
            if isinstance(target_widget, QLineEdit) and target_widget.text().strip() != target_column_name:
                target_widget.setText(target_column_name)
            if not target_column_name:
                return None, "Every mapping needs a target column name."

            dtype_widget = self.map_table.cellWidget(row, 3)
            dtype_text = (
                dtype_widget.currentText()
                if isinstance(dtype_widget, QComboBox)
                else "NVARCHAR"
            )
            java_type = import_target_data_type_to_java(dtype_text)

            length_item = self.map_table.item(row, 4)
            length_raw = length_item.text().strip() if length_item else ""

            nullable_wrap = self.map_table.cellWidget(row, 5)
            nullable_box = (
                nullable_wrap.findChild(QCheckBox)
                if isinstance(nullable_wrap, QWidget)
                else None
            )
            nullable = bool(nullable_box.isChecked()) if nullable_box is not None else True

            pk_wrap = self.map_table.cellWidget(row, 6)
            pk_label = (
                pk_wrap.findChild(QLabel) if isinstance(pk_wrap, QWidget) else None
            )
            primary_key = bool(pk_label.property("isPrimaryKey")) if pk_label else False

            mapping: dict[str, Any] = {
                "sourceColumnId": source_id_int,
                "targetColumnName": target_column_name,
                "targetDataType": java_type,
                "length": None,
                "precision": None,
                "scale": None,
                "nullable": nullable,
                "primaryKey": primary_key,
                "selected": selected,
            }
            if length_raw and length_raw != "-" and java_type in ("STRING", "DECIMAL"):
                try:
                    mapping["length"] = int(length_raw)
                except ValueError:
                    pass
            mappings.append(mapping)

        if not mappings:
            return None, "No column mappings to save."
        return (
            {
                "sheetId": sheet_id_int,
                "targetTableName": target_table,
                "targetConnectionId": target_connection_id,
                "mappings": mappings,
            },
            None,
        )

    def _on_primary(self) -> None:
        if self._busy or self._importing:
            return
        if self._step_index == 0:
            if self._remote_session and self._session_id and self._sheets:
                self._open_map_from_loaded_session()
                return
            if not self._selected_path:
                self._show_error("Select a file to upload.")
                return
            if not self._validate_upload_form():
                return
            self._go_to_map_columns()
            return
        if self._step_index == 1:
            self._go_to_execute()
            return
        if self.primary_btn.text() == "Done":
            self._on_back_to_imports()
            return
        self._start_import()

    def _on_secondary(self) -> None:
        if self._busy:
            return
        if self._step_index == 0:
            self.reset_for_new_import()
            self.exit_to_list_requested.emit()
            return
        if self._step_index == 2 and self._importing:
            self._cancel_import()
            return
        if self._step_index == 2:
            self._reset_import_progress()
            self._execute_stack.setCurrentIndex(0)
        self._clear_message()
        self._set_step(self._step_index - 1)

    def _validate_upload_form(self) -> bool:
        if not self.import_name_edit.text().strip():
            self._show_error("Import name is required.")
            return False
        if not isinstance(self.connection_combo.currentData(), dict):
            self._show_error("Select a target database connection.")
            return False
        if not self.target_table_edit.text().strip():
            self._show_error("Target table name is required.")
            return False
        size = Path(self._selected_path).stat().st_size if Path(self._selected_path).is_file() else 0
        if size > 10 * 1024 * 1024:
            self._show_error("File exceeds the 10MB limit.")
            return False
        return True

    def _show_error(self, message: str) -> None:
        show_page_error(self, self.message_label, message)

    def _clear_message(self) -> None:
        clear_page_message(self, self.message_label)
